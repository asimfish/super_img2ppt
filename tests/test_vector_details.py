import copy
import shutil
from xml.etree import ElementTree as ET

import pytest
from PIL import Image
from pptx import Presentation
from super_img2ppt.export import write_pptx, write_svg
from super_img2ppt.layout import layout_scene
from super_img2ppt.qa import preflight
from super_img2ppt.render import make_renderer, rasterize_pdf
from super_img2ppt.scene import InputError, validate_scene


def concave_scene(scene):
    page = scene["slides"][0]
    page.update(width=400, height=300)
    page["elements"] = [
        {
            "id": "icon",
            "kind": "shape",
            "shape": "freeform",
            "z": 0,
            "box": [50, 50, 120, 120],
            "vertices": [[0, 0], [1, 0], [1, 0.3], [0.3, 0.3], [0.3, 1], [0, 1]],
            "fill": "#4295FF",
        }
    ]
    return scene


def test_concave_freeform_checks_recess_and_real_solid_parts(scene, fonts, tmp_path):
    spec = validate_scene(concave_scene(scene))
    page = spec["slides"][0]
    page["elements"].append(
        {
            "id": "recess",
            "kind": "shape",
            "shape": "rect",
            "z": 1,
            "box": [110, 110, 15, 15],
            "fill": "#111111",
        }
    )
    assert preflight(spec, layout_scene(spec, fonts), tmp_path)["status"] == "pass"
    page["elements"][-1]["box"] = [60, 60, 15, 15]
    assert preflight(spec, layout_scene(spec, fonts), tmp_path)["status"] == "fail"
    page["elements"][-1].update(box=[110, 110, 15, 15], container="icon")
    assert any(
        f["code"] == "outside_container"
        for f in preflight(spec, layout_scene(spec, fonts), tmp_path)["findings"]
    )


def test_freeform_rejects_crossed_edges_and_duplicate_vertices(scene):
    spec = concave_scene(scene)
    for vertices in [
        [[0, 0], [1, 1], [0, 1], [1, 0]],
        [[0, 0], [1, 0], [1, 0], [0, 1]],
        [[0, 0], [0.5, 0], [1, 0]],
    ]:
        bad = copy.deepcopy(spec)
        bad["slides"][0]["elements"][0]["vertices"] = vertices
        with pytest.raises(InputError):
            validate_scene(bad)


def test_alpha_and_zero_width_export_as_explicit_native_properties(scene, fonts, tmp_path):
    spec = concave_scene(scene)
    icon = spec["slides"][0]["elements"][0]
    icon.update(fill_opacity=0.5, stroke="#FF0000", stroke_width=0)
    validate_scene(spec)
    pptx = tmp_path / "details.pptx"
    layouts = layout_scene(spec, fonts)
    write_pptx(spec, layouts, tmp_path, pptx)
    shape = Presentation(pptx).slides[0].shapes[0]
    assert shape._element.xpath(".//a:custGeom/a:pathLst/a:path")
    assert shape._element.xpath("./p:spPr/a:solidFill/a:srgbClr/a:alpha")[0].get("val") == "50000"
    assert shape._element.xpath("./p:spPr/a:ln/a:noFill")
    svg = tmp_path / "details.svg"
    write_svg(spec["slides"][0], layouts, tmp_path, svg)
    node = ET.parse(svg).find(".//{http://www.w3.org/2000/svg}polygon")
    assert node.get("fill-opacity") == "0.5"
    assert node.get("stroke") == "none"


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_actual_alpha_is_composited_over_underlying_color(scene, fonts, tmp_path):
    spec = concave_scene(scene)
    page = spec["slides"][0]
    page["background"] = "#204060"
    page["elements"][0].update(fill="#E0A020", fill_opacity=0.5)
    validate_scene(spec)
    pptx = tmp_path / "alpha.pptx"
    write_pptx(spec, layout_scene(spec, fonts), tmp_path, pptx)
    pdf = make_renderer().render(pptx, tmp_path / "render")
    png = rasterize_pdf(pdf, tmp_path / "render", width=400)[0]
    with Image.open(png) as im:
        assert (
            max(abs(a - b) for a, b in zip(im.getpixel((70, 70))[:3], (128, 112, 64), strict=True))
            <= 2
        )
        assert im.getpixel((130, 130))[:3] == (32, 64, 96)


def test_opacity_validation_rejects_irrelevant_and_out_of_range_fields(scene):
    spec = concave_scene(scene)
    for changes in [
        {"fill_opacity": 1.1},
        {"fill_opacity": True},
        {"fill": None, "fill_opacity": 0.5},
        {"stroke_opacity": 0.5},
        {"stroke": "#222222", "stroke_width": 0, "stroke_opacity": 0.5},
    ]:
        bad = copy.deepcopy(spec)
        bad["slides"][0]["elements"][0].update(changes)
        with pytest.raises(InputError):
            validate_scene(bad)


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_actual_gradient_and_stroke_alpha_remain_independent(scene, fonts, tmp_path):
    spec = concave_scene(scene)
    page = spec["slides"][0]
    page["elements"] = [
        {
            "id": "bar",
            "kind": "shape",
            "shape": "rect",
            "z": 0,
            "box": [40, 40, 300, 120],
            "gradient": {
                "direction": "horizontal",
                "stops": [
                    {"offset": 0, "color": "#FF0000", "opacity": 1},
                    {"offset": 1, "color": "#FF0000", "opacity": 0},
                ],
            },
            "fill_opacity": 0.5,
        },
        {
            "id": "stroke",
            "kind": "line",
            "z": 1,
            "points": [[40, 220], [340, 220]],
            "stroke": "#0000FF",
            "stroke_width": 6,
            "stroke_opacity": 0.5,
        },
    ]
    validate_scene(spec)
    pptx = tmp_path / "gradient-alpha.pptx"
    write_pptx(spec, layout_scene(spec, fonts), tmp_path, pptx)
    shapes = Presentation(pptx).slides[0].shapes
    alpha = shapes[0]._element.xpath(".//a:gs/a:srgbClr/a:alpha")
    assert [n.get("val") for n in alpha] == ["50000", "0"]
    svg = tmp_path / "gradient-alpha.svg"
    write_svg(page, layout_scene(spec, fonts), tmp_path, svg)
    assert [
        n.get("stop-opacity") for n in ET.parse(svg).findall(".//{http://www.w3.org/2000/svg}stop")
    ] == ["1", "0"]
    pdf = make_renderer().render(pptx, tmp_path / "render")
    png = rasterize_pdf(pdf, tmp_path / "render", width=400)[0]
    with Image.open(png) as im:
        for point, expected in [((190, 100), (255, 191, 191)), ((190, 220), (128, 128, 255))]:
            assert (
                max(abs(a - b) for a, b in zip(im.getpixel(point)[:3], expected, strict=True)) <= 3
            )


def test_fully_transparent_shapes_do_not_invent_collisions(scene, fonts, tmp_path):
    spec = concave_scene(scene)
    page = spec["slides"][0]
    page["elements"][0]["fill_opacity"] = 0
    page["elements"].append(
        {
            "id": "visible",
            "kind": "shape",
            "shape": "rect",
            "z": 1,
            "box": [60, 60, 15, 15],
            "fill": "#222222",
        }
    )
    assert preflight(spec, layout_scene(spec, fonts), tmp_path)["status"] == "pass"
    page["elements"][0]["fill_opacity"] = 0.1
    assert preflight(spec, layout_scene(spec, fonts), tmp_path)["status"] == "fail"
