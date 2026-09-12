import copy
import shutil
from xml.etree import ElementTree as ET

import pypdfium2 as pdfium
import pytest
from pptx import Presentation
from super_img2ppt.export import write_pptx, write_svg
from super_img2ppt.layout import layout_scene
from super_img2ppt.qa import preflight
from super_img2ppt.render import make_renderer, verify_rendered_text
from super_img2ppt.scene import InputError, validate_scene


def diagonal(scene):
    scene["slides"][0].update(
        width=640,
        height=360,
        elements=[
            {
                "id": "label",
                "kind": "text",
                "z": 2,
                "box": [180, 130, 190, 36],
                "text": "Attention tokens",
                "font_family": "Arial",
                "font_size": 22,
                "rotation": -45,
                "padding": [2, 4, 2, 4],
            }
        ],
    )
    return scene


def test_diagonal_text_is_native_and_schema_preserves_angle(scene, fonts, tmp_path):
    spec = diagonal(scene)
    validate_scene(spec)
    layouts = layout_scene(spec, fonts)
    assert preflight(spec, layouts, tmp_path)["status"] == "pass"
    pptx = tmp_path / "diagonal.pptx"
    write_pptx(spec, layouts, tmp_path, pptx)
    shape = Presentation(pptx).slides[0].shapes[0]
    assert shape.has_text_frame and shape.text == "Attention tokens"
    assert shape.rotation == 315
    svg = tmp_path / "diagonal.svg"
    write_svg(spec["slides"][0], layouts, tmp_path, svg)
    group = ET.parse(svg).find('.//{http://www.w3.org/2000/svg}g[@id="label"]')
    assert group.get("transform") == "rotate(-45 275.0 148.0)"
    for angle in [float("nan"), float("inf"), 361, True]:
        invalid = copy.deepcopy(spec)
        invalid["slides"][0]["elements"][0]["rotation"] = angle
        with pytest.raises(InputError):
            validate_scene(invalid)


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_actual_diagonal_orientation_text_and_font_checks(scene, fonts, tmp_path):
    spec = diagonal(scene)
    layouts = layout_scene(spec, fonts)
    pptx = tmp_path / "diagonal.pptx"
    write_pptx(spec, layouts, tmp_path, pptx)
    pdf = make_renderer().render(pptx, tmp_path / "render")
    report = verify_rendered_text(pdf, spec, layouts)
    assert report["status"] == "review", report
    assert report["text_boxes"][0]["matches"] is True
    assert report["text_boxes"][0]["width_measurement_axis"] == "page_x"
    assert {x["code"] for x in report["findings"]} == {
        "diagonal_glyph_bounds_require_visual_review"
    }
    with pdfium.PdfDocument(pdf) as doc:
        page = doc[0]
        tp = page.get_textpage()
        glyphs = []
        for i in range(tp.count_chars()):
            if tp.get_text_range(i, 1).strip():
                left, bottom, right, top = tp.get_charbox(i)
                glyphs.append(((left + right) / 2, (bottom + top) / 2))
        assert len(glyphs) >= 14
        dx = glyphs[-1][0] - glyphs[0][0]
        dy = glyphs[-1][1] - glyphs[0][1]
        assert dx > 30 and 0.9 < dy / dx < 1.1
        tp.close()
        page.close()


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_neighbor_inside_aabb_cannot_replace_missing_diagonal_text(scene, fonts, tmp_path):
    expected = diagonal(scene)
    expected["slides"][0]["elements"][0].update(box=[200, 100, 140, 30], text="X")
    actual = copy.deepcopy(expected)
    actual["slides"][0]["elements"][0].update(box=[215, 57, 30, 30], rotation=0)
    pptx = tmp_path / "decoy.pptx"
    write_pptx(actual, layout_scene(actual, fonts), tmp_path, pptx)
    pdf = make_renderer().render(pptx, tmp_path / "render")
    report = verify_rendered_text(pdf, expected, layout_scene(expected, fonts))
    assert report["status"] == "fail"
    assert report["text_boxes"][0]["rendered"].strip() == "X"
    assert report["text_boxes"][0]["rendered_normalized"].strip() == ""
    assert report["text_boxes"][0]["matches"] is False
    assert any(f["code"] == "rendered_text_mismatch" for f in report["findings"])
