import copy
import shutil
import xml.etree.ElementTree as ET

import pypdfium2 as pdfium
import pytest
from PIL import ImageChops
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
from super_img2ppt.export import write_pptx, write_svg
from super_img2ppt.layout import layout_scene
from super_img2ppt.qa import editability_manifest, inspect_pptx
from super_img2ppt.render import make_renderer, verify_rendered_text
from super_img2ppt.scene import InputError, validate_scene


def example(fonts):
    family = fonts.resolve("Label", None, None, False, False).family
    return {
        "version": 1,
        "slides": [
            {
                "id": "page_001",
                "width": 640,
                "height": 360,
                "reviewed": True,
                "elements": [
                    {
                        "id": "panel",
                        "kind": "shape",
                        "z": 0,
                        "box": [40, 40, 220, 160],
                        "shape": "rect",
                        "fill": "#DBEAFE",
                    },
                    {
                        "id": "label",
                        "kind": "text",
                        "z": 1,
                        "box": [60, 60, 140, 35],
                        "font_size": 20,
                        "font_family": family,
                        "text": "Encoder",
                    },
                    {
                        "id": "x",
                        "kind": "text",
                        "z": 2,
                        "box": [80, 115, 35, 40],
                        "font_size": 24,
                        "font_family": family,
                        "text": "x",
                        "rotation": -25,
                    },
                    {
                        "id": "i",
                        "kind": "text",
                        "z": 3,
                        "box": [102, 96, 25, 25],
                        "font_size": 14,
                        "font_family": family,
                        "text": "i",
                        "rotation": -25,
                    },
                    {
                        "id": "neighbor",
                        "kind": "text",
                        "z": 4,
                        "box": [380, 100, 200, 40],
                        "font_size": 20,
                        "font_family": family,
                        "text": "Unchanged",
                    },
                ],
                "groups": [
                    {
                        "id": "encoder",
                        "members": ["formula", "panel", "label"],
                        "description": "Move the complete encoder module",
                    },
                    {
                        "id": "formula",
                        "members": ["i", "x"],
                        "description": "Keep variable and superscript together",
                    },
                ],
            }
        ],
    }


def test_nested_groups_preserve_native_leaves_and_svg_order(fonts, tmp_path):
    scene = validate_scene(example(fonts))
    manifest = editability_manifest(scene)["slides"][0]
    assert manifest["native_leaf_objects"] == 5
    assert manifest["top_level_selectable_objects"] == 2
    assert manifest["root_groups"] == ["encoder"]
    assert manifest["ungrouped_elements"] == ["neighbor"]
    layouts = layout_scene(scene, fonts)
    write_pptx(scene, layouts, tmp_path, tmp_path / "grouped.pptx")
    write_svg(scene["slides"][0], layouts, tmp_path, tmp_path / "grouped.svg")
    deck = Presentation(tmp_path / "grouped.pptx")
    shapes = deck.slides[0].shapes
    assert [s.name for s in shapes] == ["encoder", "neighbor"]
    assert shapes[0].shape_type == MSO_SHAPE_TYPE.GROUP
    assert [s.name for s in shapes[0].shapes] == ["panel", "label", "formula"]
    assert [s.name for s in shapes[0].shapes[2].shapes] == ["x", "i"]
    assert inspect_pptx(tmp_path / "grouped.pptx", scene)["status"] == "pass"
    svg = ET.parse(tmp_path / "grouped.svg").getroot()
    ns = {"s": "http://www.w3.org/2000/svg"}
    assert [g.get("id") for g in svg.findall("s:g", ns)] == ["encoder", "neighbor"]
    assert [g.get("id") for g in svg.find("s:g", ns).findall("s:g", ns)] == [
        "panel",
        "label",
        "formula",
    ]
    # A missing nested text object must not be hidden by accepting its outer group.
    child = shapes[0].shapes[2].shapes[0]
    child._element.getparent().remove(child._element)
    deck.save(tmp_path / "missing.pptx")
    assert inspect_pptx(tmp_path / "missing.pptx", scene)["status"] == "fail"


@pytest.mark.parametrize(
    "change",
    [
        lambda g: g.append({"id": "other", "members": ["x", "neighbor"]}),
        lambda g: g[1].update(members=["x", "missing"]),
        lambda g: g[1].update(members=["x", "x"]),
        lambda g: g[1].update(members=["x", "encoder"]),
        lambda g: g[1].update(id="label"),
        lambda g: g[1].update(members=["x"]),
        lambda g: g[0].update(members=["formula", "panel", "neighbor"]),
    ],
)
def test_ambiguous_groups_are_rejected(fonts, change):
    scene = example(fonts)
    change(scene["slides"][0]["groups"])
    with pytest.raises(InputError):
        validate_scene(scene)


def pixels(pdf):
    with pdfium.PdfDocument(pdf) as doc:
        page = doc[0]
        bitmap = page.render(scale=1)
        image = bitmap.to_pil().convert("RGB")
        bitmap.close()
        page.close()
        return image


def pdf_text_centers(pdf):
    with pdfium.PdfDocument(pdf) as doc:
        page = doc[0]
        tp = page.get_textpage()
        result = []
        for i in range(tp.count_chars()):
            char = chr(pdfium.raw.FPDFText_GetUnicode(tp, i))
            if not char.isspace():
                x0, y0, x1, y1 = tp.get_charbox(i)
                result.append((char, (x0 + x1) / 2, (y0 + y1) / 2))
        tp.close()
        page.close()
        return result


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_actual_groups_preserve_pixels_and_support_module_move_and_label_edit(fonts, tmp_path):
    scene = validate_scene(example(fonts))
    flat = copy.deepcopy(scene)
    flat["slides"][0].pop("groups")
    layouts = layout_scene(scene, fonts)
    renderer = make_renderer()
    pdfs = {}
    for name, spec in [("flat", flat), ("grouped", scene)]:
        path = tmp_path / (name + ".pptx")
        write_pptx(spec, layouts, tmp_path, path)
        pdfs[name] = renderer.render(path, tmp_path / name)
    assert ImageChops.difference(pixels(pdfs["flat"]), pixels(pdfs["grouped"])).getbbox() is None
    assert verify_rendered_text(pdfs["grouped"], scene, layouts)["status"] in {"pass", "review"}
    deck = Presentation(tmp_path / "grouped.pptx")
    outer = deck.slides[0].shapes[0]
    outer.left += Inches(0.5)
    deck.save(tmp_path / "moved.pptx")
    moved = renderer.render(tmp_path / "moved.pptx", tmp_path / "moved")
    before, after = pdf_text_centers(pdfs["grouped"]), pdf_text_centers(moved)
    assert len(before) == len(after)
    for old, new in zip(before, after, strict=True):
        assert old[0] == new[0]
        assert new[1] - old[1] == pytest.approx(36 if old[1] < 450 else 0, abs=0.1)
        assert new[2] == pytest.approx(old[2], abs=0.1)
    # Edit one child through its retained text run; sibling formula and neighbor stay intact.
    label = outer.shapes[1]
    label.text_frame.paragraphs[0].runs[0].text = "Decoder"
    deck.save(tmp_path / "edited.pptx")
    check = Presentation(tmp_path / "edited.pptx").slides[0].shapes
    assert check[0].shapes[1].text == "Decoder"
    assert check[0].shapes[2].shapes[0].text == "x"
    assert check[1].text == "Unchanged"


def test_group_nesting_bound_is_checked_from_roots(fonts):
    scene = example(fonts)
    slide = scene["slides"][0]
    template = slide["elements"][0]
    slide["elements"] = [{**template, "id": f"leaf{i}", "z": i} for i in range(10)]
    slide["groups"] = [{"id": "g0", "members": ["leaf0", "leaf1"]}]
    for i in range(1, 8):
        slide["groups"].append({"id": f"g{i}", "members": [f"g{i - 1}", f"leaf{i + 1}"]})
    validate_scene(scene)
    slide["groups"].append({"id": "g8", "members": ["g7", "leaf9"]})
    with pytest.raises(InputError, match="eight"):
        validate_scene(scene)


@pytest.mark.render
@pytest.mark.parametrize("vertical", [False, True])
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_collinear_segment_groups_keep_the_actual_stroke(tmp_path, vertical):
    points = [[40, 70], [100, 70], [160, 70]]
    if vertical:
        points = [[y, x] for x, y in points]
    elements = [
        {
            "id": f"line{i}",
            "kind": "line",
            "z": i,
            "points": [points[i], points[i + 1]],
            "stroke": "#0000FF",
            "stroke_width": 2,
        }
        for i in range(2)
    ]
    scene = {
        "version": 1,
        "slides": [{"id": "page_001", "width": 240, "height": 240, "elements": elements}],
    }
    renderer = make_renderer()
    write_pptx(scene, {}, tmp_path, tmp_path / "flat.pptx")
    flat = renderer.render(tmp_path / "flat.pptx", tmp_path / "flat")
    scene["slides"][0]["groups"] = [{"id": "curve", "members": ["line0", "line1"]}]
    validate_scene(scene)
    write_pptx(scene, {}, tmp_path, tmp_path / "grouped.pptx")
    grouped = renderer.render(tmp_path / "grouped.pptx", tmp_path / "grouped")
    assert ImageChops.difference(pixels(flat), pixels(grouped)).getbbox() is None
