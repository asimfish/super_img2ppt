import copy
import shutil
from xml.etree import ElementTree as ET

import pytest
from PIL import Image
from pptx import Presentation
from super_img2ppt.export import write_pptx, write_svg
from super_img2ppt.layout import layout_scene
from super_img2ppt.render import make_renderer, rasterize_pdf
from super_img2ppt.scene import InputError, validate_scene


def colorbar(scene, direction="vertical"):
    scene["slides"][0].update(width=400, height=400)
    scene["slides"][0]["elements"] = [
        {
            "id": "bar",
            "kind": "shape",
            "shape": "rect",
            "box": [50, 50, 300, 300],
            "z": 0,
            "gradient": {
                "direction": direction,
                "stops": [
                    {"offset": 0, "color": "#FFFFE5"},
                    {"offset": 0.5, "color": "#78C679"},
                    {"offset": 1, "color": "#004529"},
                ],
            },
        },
    ]
    return scene


def test_gradient_contract_rejects_ambiguous_or_unbounded_stops(scene):
    spec = colorbar(scene)
    validate_scene(spec)
    for stops in [
        [],
        [{"offset": 0, "color": "#000000"}],
        [{"offset": 0, "color": "#000000"}, {"offset": 0, "color": "#FFFFFF"}],
        [{"offset": 0.2, "color": "#000000"}, {"offset": 1, "color": "#FFFFFF"}],
        [{"offset": i / 16, "color": "#000000"} for i in range(17)],
    ]:
        bad = copy.deepcopy(spec)
        bad["slides"][0]["elements"][0]["gradient"]["stops"] = stops
        with pytest.raises(InputError):
            validate_scene(bad)
    spec["slides"][0]["elements"][0]["fill"] = "#FF0000"
    with pytest.raises(InputError):
        validate_scene(spec)


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
@pytest.mark.parametrize("direction", ["vertical", "horizontal"])
def test_single_native_gradient_has_continuous_actual_colorbar(scene, fonts, tmp_path, direction):
    spec = colorbar(scene, direction)
    validate_scene(spec)
    layouts = layout_scene(spec, fonts)
    pptx = tmp_path / "gradient.pptx"
    write_pptx(spec, layouts, tmp_path, pptx)
    shapes = Presentation(pptx).slides[0].shapes
    assert len(shapes) == 1
    stops = shapes[0]._element.xpath(".//a:gradFill/a:gsLst/a:gs")
    assert [s.get("pos") for s in stops] == ["0", "50000", "100000"]
    svg = tmp_path / "gradient.svg"
    write_svg(spec["slides"][0], layouts, tmp_path, svg)
    assert len(ET.parse(svg).findall(".//{http://www.w3.org/2000/svg}stop")) == 3
    pdf = make_renderer().render(pptx, tmp_path / "render")
    png = rasterize_pdf(pdf, tmp_path / "render", width=400)[0]
    with Image.open(png) as im:
        pixels = [
            im.getpixel((200, v) if direction == "vertical" else (v, 200))[:3]
            for v in range(53, 347)
        ]
    assert pixels[0][0] > 245 and pixels[0][1] > 245
    assert pixels[-1][0] < 10 and 65 < pixels[-1][1] < 80
    assert all(
        max(abs(a - b) for a, b in zip(p, q, strict=True)) <= 5
        for p, q in zip(pixels, pixels[1:], strict=False)
    )
