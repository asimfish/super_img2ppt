import copy
import shutil
from xml.etree import ElementTree as ET

import pytest
from PIL import Image
from pptx import Presentation
from super_img2ppt.export import write_pptx, write_svg
from super_img2ppt.layout import layout_scene
from super_img2ppt.qa import preflight
from super_img2ppt.render import make_renderer, rasterize_pdf
from super_img2ppt.scene import InputError, validate_scene


def trapezoid(scene):
    page = scene["slides"][0]
    page.update(width=480, height=260)
    page["elements"] = [
        {
            "id": "encoder",
            "kind": "shape",
            "shape": "polygon",
            "z": 0,
            "box": [100, 30, 200, 180],
            "vertices": [[0, 0], [1, 0.2], [1, 0.8], [0, 1]],
            "fill": "#C080D0",
            "stroke": "#000000",
            "stroke_width": 2,
        },
        {
            "id": "label",
            "kind": "text",
            "z": 1,
            "box": [120, 85, 140, 40],
            "text": "Encoder",
            "font_family": "Arial",
            "font_size": 24,
            "container": "encoder",
        },
    ]
    return scene


def test_convex_polygon_contract_and_rejected_ambiguous_shapes(scene):
    spec = trapezoid(scene)
    assert validate_scene(spec) == spec
    bad_vertices = [
        [[0, 0], [1, 1], [0, 1], [1, 0]],
        [[0, 0], [1, 0], [0.4, 0.4], [1, 1], [0, 1]],
        [[0, 0], [1, 0], [1, 0], [0, 1]],
        [[0, 0], [0.5, 0], [1, 0]],
        [[0, 0], [2, 0], [1, 1]],
        [[0.5, 0], [0.79, 0.9], [0, 0.35], [1, 0.35], [0.21, 0.9]],
    ]
    for vertices in bad_vertices:
        invalid = copy.deepcopy(spec)
        invalid["slides"][0]["elements"][0]["vertices"] = vertices
        with pytest.raises(InputError):
            validate_scene(invalid)
    for change in [{"shape": "rect"}, {"vertices": None}]:
        invalid = copy.deepcopy(spec)
        invalid["slides"][0]["elements"][0].update(change)
        with pytest.raises(InputError):
            validate_scene(invalid)


@pytest.mark.parametrize("reverse", [False, True])
def test_polygon_containment_and_empty_corner_use_visible_geometry(scene, fonts, tmp_path, reverse):
    spec = trapezoid(scene)
    page = spec["slides"][0]
    if reverse:
        page["elements"][0]["vertices"].reverse()
    validate_scene(spec)
    page["elements"].append(
        {
            "id": "corner",
            "kind": "shape",
            "shape": "rect",
            "z": 2,
            "box": [270, 34, 10, 10],
            "fill": "#FF0000",
        }
    )
    assert preflight(spec, layout_scene(spec, fonts), tmp_path)["status"] == "pass"
    page["elements"][-1]["box"][1] = 80
    report = preflight(spec, layout_scene(spec, fonts), tmp_path)
    assert any(f["code"] == "unintended_overlap" for f in report["findings"])
    page["elements"].pop()
    page["elements"][1]["box"] = [260, 30, 140, 40]
    report = preflight(spec, layout_scene(spec, fonts), tmp_path)
    assert any(f["code"] == "outside_container" for f in report["findings"])


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_polygon_is_editable_and_actual_render_keeps_slanted_corners(scene, fonts, tmp_path):
    spec = trapezoid(scene)
    validate_scene(spec)
    layouts = layout_scene(spec, fonts)
    pptx = tmp_path / "polygon.pptx"
    write_pptx(spec, layouts, tmp_path, pptx)
    deck = Presentation(pptx)
    assert len(deck.slides[0].shapes) == 2
    encoder = deck.slides[0].shapes[0]
    assert encoder._element.xpath(".//a:custGeom/a:pathLst/a:path")
    assert deck.slides[0].shapes[1].text == "Encoder"
    svg = tmp_path / "polygon.svg"
    write_svg(spec["slides"][0], layouts, tmp_path, svg)
    polygon = ET.parse(svg).find(".//{http://www.w3.org/2000/svg}polygon")
    points = [tuple(map(float, p.split(","))) for p in polygon.get("points").split()]
    assert set(points) == {(100, 30), (300, 66), (300, 174), (100, 210)}
    pdf = make_renderer().render(pptx, tmp_path / "render")
    png = rasterize_pdf(pdf, tmp_path / "render", width=480)[0]
    with Image.open(png) as im:
        assert min(im.getpixel((275, 40))[:3]) > 245
        assert min(im.getpixel((275, 195))[:3]) > 245
        for xy in [(110, 45), (285, 85), (285, 155), (110, 190)]:
            r, g, b = im.getpixel(xy)[:3]
            assert r > 160 and b > 160 and g < 155
