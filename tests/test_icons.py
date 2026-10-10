import shutil

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from super_img2ppt.export import write_pptx
from super_img2ppt.icons import trace_icon, trace_icon_file
from super_img2ppt.layout import layout_scene
from super_img2ppt.render import make_renderer, rasterize_pdf
from super_img2ppt.scene import InputError


def icon_image():
    image = Image.new("RGB", (80, 80), "white")
    ImageDraw.Draw(image).polygon(
        [(12, 12), (60, 12), (60, 24), (24, 24), (24, 60), (12, 60)], fill="#4295FF"
    )
    return image


def test_trace_one_concave_icon_is_a_single_native_object():
    result = trace_icon(icon_image(), [5, 5, 65, 65], "#4295FF")
    assert result["status"] == "review"
    assert len(result["elements"]) == 1
    assert result["vertex_count"] == 6
    assert result["elements"][0]["shape"] == "freeform"
    assert result["mask_pixels"] == 1105


@pytest.mark.parametrize("problem", ["hole", "separate", "corner", "edge", "empty"])
def test_icon_refuses_ambiguous_topology_instead_of_losing_parts(problem):
    image = icon_image()
    draw = ImageDraw.Draw(image)
    if problem == "hole":
        draw.rectangle((16, 16, 19, 19), fill="white")
    elif problem == "separate":
        draw.rectangle((45, 45, 50, 50), fill="#4295FF")
    elif problem == "corner":
        draw.rectangle((25, 61, 26, 62), fill="#4295FF")
    elif problem == "edge":
        draw.rectangle((5, 12, 12, 16), fill="#4295FF")
    else:
        image = Image.new("RGB", (80, 80), "white")
    with pytest.raises(InputError):
        trace_icon(image, [5, 5, 65, 65], "#4295FF")


def test_fully_transparent_pixels_are_not_traced():
    image = Image.new("RGBA", (80, 80), (66, 149, 255, 0))
    with pytest.raises(InputError, match="empty"):
        trace_icon(image, [5, 5, 65, 65], "#4295FF")


def test_icon_file_emits_hash_and_source_coordinate_geometry(tmp_path):
    source = tmp_path / "source.png"
    icon_image().save(source)
    result = trace_icon_file(source, tmp_path / "trace", roi=[5, 5, 65, 65], color="#4295FF")
    assert len(result["source_sha256"]) == 64
    assert result["elements"][0]["box"] == [12, 12, 49, 49]
    assert (tmp_path / "trace" / "overlay.png").is_file()


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_actual_native_icon_retains_color_recess_and_editable_points(fonts, tmp_path):
    image = icon_image()
    result = trace_icon(image, [5, 5, 65, 65], "#4295FF")
    spec = {
        "version": 1,
        "slides": [{"id": "page_001", "width": 80, "height": 80, "elements": result["elements"]}],
    }
    pptx = tmp_path / "icon.pptx"
    write_pptx(spec, layout_scene(spec, fonts), tmp_path, pptx)
    shape = Presentation(pptx).slides[0].shapes[0]
    assert shape._element.xpath(".//a:custGeom/a:pathLst/a:path/a:lnTo")
    assert not shape._element.xpath(".//a:blip")
    pdf = make_renderer().render(pptx, tmp_path / "render")
    png = rasterize_pdf(pdf, tmp_path / "render", width=80)[0]
    with Image.open(png) as actual:
        assert actual.getpixel((18, 40))[:3] == (66, 149, 255)
        assert actual.getpixel((40, 40))[:3] == (255, 255, 255)


def test_constant_alpha_is_retained_and_varying_alpha_is_rejected():
    image = icon_image().convert("RGBA")
    image.putalpha(128)
    result = trace_icon(image, [5, 5, 65, 65], "#4295FF")
    assert result["elements"][0]["fill_opacity"] == pytest.approx(128 / 255)
    image.putpixel((18, 18), (66, 149, 255, 64))
    with pytest.raises(InputError, match="varying foreground alpha"):
        trace_icon(image, [5, 5, 65, 65], "#4295FF")


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_real_gr00t_icon_render_preserves_source_contour(fonts, tmp_path):
    from pathlib import Path

    from super_img2ppt.regions import compare_region

    source = Image.open(Path(__file__).parent / "fixtures/icons/gr00t_snowflake.png").convert("RGB")
    result = trace_icon(source, [0, 0, 59, 60], "#4295FF", tolerance=64, simplify_px=0.75)
    spec = {
        "version": 1,
        "slides": [{"id": "page_001", "width": 59, "height": 60, "elements": result["elements"]}],
    }
    pptx = tmp_path / "real-icon.pptx"
    write_pptx(spec, layout_scene(spec, fonts), tmp_path, pptx)
    pdf = make_renderer().render(pptx, tmp_path / "render")
    png = rasterize_pdf(pdf, tmp_path / "render", width=59)[0]
    with Image.open(png) as raw:
        actual = raw.convert("RGB").crop((0, 0, 59, 60))
    metrics = compare_region(source, actual, [0, 0, 59, 60], "#4295FF", tolerance=64)
    assert metrics["mask_iou"] >= 0.9
    assert all(abs(v) <= 1 for v in metrics["edge_delta_px"])
    assert len(Presentation(pptx).slides[0].shapes) == 1
