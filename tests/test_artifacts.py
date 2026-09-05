import json
import shutil
import subprocess
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from PIL import Image
from pptx import Presentation
from super_img2ppt.cli import build
from super_img2ppt.cli import main as cli_main
from super_img2ppt.export import write_pptx, write_svg
from super_img2ppt.layout import layout_scene
from super_img2ppt.prepare import prepare
from super_img2ppt.qa import inspect_pptx
from super_img2ppt.render import make_renderer, verify_rendered_text
from super_img2ppt.scene import load_scene

RENDER_AVAILABLE = shutil.which("soffice") is not None


def test_native_editability_fonts_no_effects_and_notes(scene, fonts, tmp_path):
    scene["slides"][0]["elements"].append(
        {
            "id": "rect",
            "kind": "shape",
            "z": 1,
            "box": [80, 250, 300, 180],
            "shape": "round_rect",
            "fill": "#DDEEFF",
        }
    )
    layouts = layout_scene(scene, fonts)
    out = tmp_path / "native.pptx"
    write_pptx(scene, layouts, tmp_path, out)
    assert inspect_pptx(out, scene)["status"] == "pass"
    pptx = Presentation(out)
    text_shape = next(s for s in pptx.slides[0].shapes if s.name == "title")
    assert text_shape.has_text_frame
    assert text_shape.text == "Hello 中文 123"
    assert text_shape.text_frame.word_wrap is False
    for run in text_shape._element.xpath(".//a:rPr"):
        assert run.find("{http://schemas.openxmlformats.org/drawingml/2006/main}latin") is not None
        assert run.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea") is not None
        assert run.find("{http://schemas.openxmlformats.org/drawingml/2006/main}cs") is not None
    rect = next(s for s in pptx.slides[0].shapes if s.name == "rect")
    assert rect.shadow.inherit is False
    assert all(node.get("idx") == "0" for node in rect._element.xpath("./p:style/a:effectRef"))
    text_shape.text = "Changed by editor"
    pptx.save(tmp_path / "edited.pptx")
    edited = Presentation(tmp_path / "edited.pptx")
    assert next(s for s in edited.slides[0].shapes if s.name == "title").text == "Changed by editor"


def test_svg_escapes_content_and_contains_native_elements(scene, fonts, tmp_path):
    scene["slides"][0]["elements"][0]["text"] = '<script>alert("x")</script> & text'
    layouts = layout_scene(scene, fonts)
    out = tmp_path / "page.svg"
    write_svg(scene["slides"][0], layouts, tmp_path, out)
    xml = ET.parse(out)
    ns = {"s": "http://www.w3.org/2000/svg"}
    assert xml.findall(".//s:text", ns)
    assert xml.findall(".//s:script", ns) == []
    assert '<script>alert("x")</script> & text' in "".join(xml.getroot().itertext())


def test_image_contain_preserves_ratio_and_can_be_moved(scene, fonts, tmp_path):
    Image.new("RGB", (200, 100), "red").save(tmp_path / "asset.png")
    scene["slides"][0]["elements"].append(
        {
            "id": "picture",
            "kind": "image",
            "z": 0,
            "box": [80, 250, 300, 300],
            "path": "asset.png",
            "contains_text": False,
            "provenance": "Original synthetic asset for testing",
        }
    )
    out = tmp_path / "image.pptx"
    write_pptx(scene, layout_scene(scene, fonts), tmp_path, out)
    pptx = Presentation(out)
    picture = next(s for s in pptx.slides[0].shapes if s.name == "picture")
    assert picture.width / picture.height == pytest.approx(2)
    assert picture.top > 250 * 914400 * 13.333333 / 1280


def test_build_refuses_bad_geometry_and_marks_unrendered_draft(scene, tmp_path):
    source = tmp_path / "scene.json"
    scene["slides"][0]["elements"][0]["box"][2] = 10
    source.write_text(json.dumps(scene))
    out = tmp_path / "bad"
    report = build(source, out, render=False)
    assert report["status"] == "fail"
    assert not (out / "editable.pptx").exists()
    assert (out / "validation.json").is_file()
    scene["slides"][0]["elements"][0]["box"][2] = 700
    source.write_text(json.dumps(scene))
    report = build(source, tmp_path / "draft", render=False)
    assert report["status"] == "unverified"
    assert report["automated_checks"]["rendered_text"]["status"] == "not_run"


@pytest.mark.parametrize("command", ["check", "build"])
@pytest.mark.parametrize("failure", ["timeout", "nonzero"])
def test_font_discovery_failure_is_recorded_without_uncaught_traceback(
    scene, tmp_path, monkeypatch, capsys, command, failure
):
    from super_img2ppt import fonts as font_module

    def fail_discovery(*args, **kwargs):
        if failure == "timeout":
            raise subprocess.TimeoutExpired("fc-list", 30)
        raise subprocess.CalledProcessError(1, "fc-list", stderr="fontconfig failed")

    source = tmp_path / "scene.json"
    source.write_text(json.dumps(scene))
    out = tmp_path / "failed"
    font_module.discover_fonts.cache_clear()
    monkeypatch.setattr(font_module.shutil, "which", lambda name: "/test/fc-list")
    monkeypatch.setattr(font_module.subprocess, "run", fail_discovery)
    try:
        assert cli_main([command, str(source), "--out", str(out)]) == 2
        report = json.loads((out / "validation.json").read_text())
        assert report["status"] == "fail"
        assert "font discovery" in report["error"]
        assert not (out / "editable.pptx").exists()
        assert "Traceback" not in capsys.readouterr().err
    finally:
        font_module.discover_fonts.cache_clear()


@pytest.mark.render
@pytest.mark.skipif(not RENDER_AVAILABLE, reason="LibreOffice not installed")
def test_actual_render_and_resolved_scene_roundtrip(tmp_path):
    root = Path(__file__).resolve().parents[1]
    report = build(root / "examples/reconstruction.json", tmp_path / "first")
    assert report["status"] in {"pass", "review"}, report
    assert report["automated_checks"]["native_objects"]["status"] == "pass"
    assert report["automated_checks"]["rendered_text"]["status"] in {"pass", "review"}
    resolved = load_scene(tmp_path / "first/scene.resolved.json")
    assert len(resolved["slides"]) == 2
    assert all(
        box["actual_fonts"] for box in report["automated_checks"]["rendered_text"]["text_boxes"]
    )
    again = build(tmp_path / "first/scene.resolved.json", tmp_path / "again")
    assert again["status"] in {"pass", "review"}, again


@pytest.mark.render
@pytest.mark.skipif(not RENDER_AVAILABLE, reason="LibreOffice not installed")
def test_render_detects_font_substitution_not_just_text(scene, fonts, tmp_path):
    element = scene["slides"][0]["elements"][0]
    element.update(text="Exact font required", font_family="Arial")
    layouts = layout_scene(scene, fonts)
    pptx = tmp_path / "tampered.pptx"
    write_pptx(scene, layouts, tmp_path, pptx)
    doc = Presentation(pptx)
    for shape in doc.slides[0].shapes:
        for node in shape._element.xpath(".//a:latin | .//a:ea | .//a:cs"):
            node.set("typeface", "Times New Roman")
    doc.save(pptx)
    pdf = make_renderer().render(pptx, tmp_path / "render")
    result = verify_rendered_text(pdf, scene, layouts)
    assert result["text_boxes"][0]["matches"]
    assert any(f["code"] == "renderer_font_substitution" for f in result["findings"])


@pytest.mark.render
@pytest.mark.skipif(not RENDER_AVAILABLE, reason="LibreOffice not installed")
def test_render_flags_spacing_drift_with_correct_text_and_font(scene, fonts, tmp_path):
    element = scene["slides"][0]["elements"][0]
    element.update(text="Spacing must remain visible", font_family="Arial", font_size=28)
    layouts = layout_scene(scene, fonts)
    pptx = tmp_path / "tracking.pptx"
    write_pptx(scene, layouts, tmp_path, pptx)
    doc = Presentation(pptx)
    for shape in doc.slides[0].shapes:
        for node in shape._element.xpath(".//a:rPr"):
            node.set("spc", "200")  # Inject 2 pt tracking without changing text or typeface.
    doc.save(pptx)
    pdf = make_renderer().render(pptx, tmp_path / "render")
    result = verify_rendered_text(pdf, scene, layouts)
    assert result["text_boxes"][0]["matches"]
    assert result["status"] == "review", result
    assert {f["code"] for f in result["findings"]} == {"rendered_ink_width_drift"}


@pytest.mark.render
@pytest.mark.skipif(not RENDER_AVAILABLE, reason="LibreOffice not installed")
def test_pdf_and_pptx_prepare_preserve_page_order_and_notes(scene, fonts, tmp_path):
    scene["slides"].append({**scene["slides"][0], "id": "page_002", "notes": "Second page note"})
    pptx = tmp_path / "input.pptx"
    write_pptx(scene, layout_scene(scene, fonts), tmp_path, pptx)
    result = prepare([pptx], tmp_path / "pptx_job", ocr="none")
    draft = load_scene(tmp_path / "pptx_job/scene.json")
    assert result["page_count"] == 2
    assert [s["notes"] for s in draft["slides"]] == ["原样保留\nsecond line", "Second page note"]
    pdf = make_renderer().render(pptx, tmp_path / "pdf_input")
    pdf_result = prepare([pdf], tmp_path / "pdf_job", ocr="none")
    assert pdf_result["page_count"] == 2
