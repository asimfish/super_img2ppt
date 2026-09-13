import hashlib
import json
import zipfile

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from super_img2ppt.native import M, replace_rasters, validate_expression
from super_img2ppt.scene import InputError


def inputs(tmp_path):
    image = tmp_path / "crop.png"
    Image.new("RGB", (100, 50), "red").save(image)
    deck = Presentation()
    deck.slide_width = Inches(10)
    deck.slide_height = Inches(5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    pic = slide.shapes.add_picture(str(image), Inches(1), Inches(1), Inches(3), Inches(1))
    pic.name = "target"
    slide.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(1)).name = "untouched"
    source = tmp_path / "original.pptx"
    deck.save(source)
    plan = dict(
        version=1,
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        width=1000,
        height=500,
        replacements=[
            dict(
                slide=1,
                image="target",
                items=[
                    dict(
                        id="block",
                        kind="cuboid",
                        box=[100, 100, 40, 60],
                        depth=[8, 6],
                        front="#B08080",
                        side="#905050",
                        top="#D0A0A0",
                    ),
                    dict(
                        id="formula",
                        kind="equation",
                        box=[160, 100, 200, 60],
                        font_size=24,
                        color="#101522",
                        expression=[
                            {"base": {"hat": "x"}, "sub": "0", "sup": "2"},
                            "=",
                            {"num": "a", "den": {"sqrt": "b"}},
                        ],
                    ),
                ],
            )
        ],
    )
    return source, plan


def test_replacement_roundtrip_native_structure_and_source_unchanged(tmp_path):
    source, plan = inputs(tmp_path)
    result = replace_rasters(source, plan, tmp_path / "out", False)
    assert result["status"] == "unverified"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == plan["source_sha256"]
    deck = Presentation(tmp_path / "out/editable.pptx")
    shapes = deck.slides[0].shapes
    assert [s.name for s in shapes] == ["block", "formula", "untouched"]
    assert len(shapes[0].shapes) == 3
    assert all(
        ref.get("idx") == "0"
        for face in shapes[0].shapes
        for ref in face._element.xpath("./p:style/a:effectRef")
    )
    assert all(
        s._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst")
        is not None
        for s in shapes[0].shapes
    )
    root = shapes._spTree
    ids = root.xpath(".//p:cNvPr/@id")
    assert len(ids) == len(set(ids))
    assert not root.xpath(".//p:pic")
    for tag in ["oMath", "sSubSup", "acc", "f", "rad"]:
        assert len(root.findall(f".//{{{M}}}{tag}")) == 1
    assert not result["remaining_images"]
    # Saving through python-pptx must retain true equation XML.
    deck.save(tmp_path / "roundtrip.pptx")
    with zipfile.ZipFile(tmp_path / "roundtrip.pptx") as archive:
        xml = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        assert len(xml.findall(f".//{{{M}}}oMath")) == 1


@pytest.mark.parametrize("change", ["hash", "target", "duplicate", "outside", "expression"])
def test_invalid_replacements_do_not_emit_deliverable(tmp_path, change):
    source, plan = inputs(tmp_path)
    if change == "hash":
        plan["source_sha256"] = "0" * 64
    if change == "target":
        plan["replacements"][0]["image"] = "untouched"
    if change == "duplicate":
        plan["replacements"].append(plan["replacements"][0])
    if change == "outside":
        plan["replacements"][0]["items"][0]["box"][0] = 1000
    if change == "expression":
        plan["replacements"][0]["items"][1]["expression"] = {"xml": "<entity/>"}
    with pytest.raises(InputError):
        replace_rasters(source, plan, tmp_path / "out", False)
    assert not (tmp_path / "out/editable.pptx").exists()
    assert json.loads((tmp_path / "out/native.json").read_text())["status"] == "fail"


def test_equation_depth_and_nodes_bounded():
    expr = "x"
    for _ in range(20):
        expr = {"hat": expr}
    with pytest.raises(InputError):
        validate_expression(expr)
    with pytest.raises(InputError):
        validate_expression([["x"] * 128] * 5)


@pytest.mark.render
def test_real_native_equation_and_blocks_render(tmp_path):
    source, plan = inputs(tmp_path)
    result = replace_rasters(source, plan, tmp_path / "out")
    assert result["status"] == "review"
    assert (tmp_path / "out/render/page_001.png").exists()
    import pypdfium2 as pdfium

    with pdfium.PdfDocument(tmp_path / "out/render/editable.pdf") as pdf:
        page = pdf[0]
        text = page.get_textpage()
        actual = text.get_text_range()
        text.close()
        page.close()
    assert all(c in actual for c in ["x", "0", "2", "a", "b"])


def test_nested_pictures_remain_disclosed_and_name_conflicts_fail(tmp_path):
    source, plan = inputs(tmp_path)
    deck = Presentation(source)
    slide = deck.slides[0]
    pic = slide.shapes.add_picture(str(tmp_path / "crop.png"), 0, 0)
    pic.name = "nested"
    slide.shapes.add_group_shape([pic])
    deck.save(source)
    plan["source_sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    report = replace_rasters(source, plan, tmp_path / "nested-out", False)
    assert {r["name"] for r in report["remaining_images"]} == {"nested"}
    plan["replacements"][0]["items"][0]["id"] = "nested"
    with pytest.raises(InputError, match="conflicts"):
        replace_rasters(source, plan, tmp_path / "conflict-out", False)
