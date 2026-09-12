import copy
import shutil
from xml.etree import ElementTree as ET

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from super_img2ppt.export import write_pptx, write_svg
from super_img2ppt.layout import layout_scene
from super_img2ppt.qa import preflight
from super_img2ppt.render import make_renderer, rasterize_pdf, verify_rendered_text
from super_img2ppt.scene import InputError, validate_scene


def diagram(scene):
    page = scene["slides"][0]
    page.update(width=640, height=360)
    page["elements"] = [
        {
            "id": "vertical",
            "kind": "text",
            "z": 2,
            "box": [90, 135, 180, 40],
            "text": "Down Sampling",
            "font_family": "Arial",
            "font_size": 22,
            "rotation": 90,
            "align": "center",
            "valign": "middle",
        },
        {
            "id": "arrow",
            "kind": "line",
            "z": 1,
            "points": [[280, 90], [480, 90]],
            "stroke": "#FF0000",
            "stroke_width": 2,
            "arrow": True,
            "arrow_head": {"length": 24, "width": 20},
        },
        {
            "id": "dashed",
            "kind": "line",
            "z": 1,
            "points": [[280, 230], [480, 230]],
            "stroke": "#0000FF",
            "stroke_width": 2,
            "dash": [12, 8],
        },
    ]
    return scene


def test_diagram_contract_rejects_silent_style_loss(scene):
    spec = diagram(scene)
    assert validate_scene(spec) == spec
    for element, changes in [
        (0, {"rotation": 361}),
        (1, {"arrow": False}),
        (1, {"arrow_head": {"length": 300, "width": 20}}),
        (2, {"dash": [0, 4]}),
        (2, {"stroke_width": 0}),
    ]:
        invalid = copy.deepcopy(spec)
        invalid["slides"][0]["elements"][element].update(changes)
        with pytest.raises(InputError):
            validate_scene(invalid)


def test_rotated_ink_and_custom_head_participate_in_collision_checks(scene, fonts, tmp_path):
    spec = diagram(scene)
    page = spec["slides"][0]
    assert preflight(spec, layout_scene(spec, fonts), tmp_path)["status"] == "pass"
    page["elements"].append(
        {
            "id": "obstacle",
            "kind": "shape",
            "z": 3,
            "shape": "rect",
            "box": [458, 81, 5, 3],
            "fill": "#000000",
        }
    )
    report = preflight(spec, layout_scene(spec, fonts), tmp_path)
    assert any(
        f["code"] == "unintended_overlap" and f["elements"] == ["arrow", "obstacle"]
        for f in report["findings"]
    )
    page["elements"][-1]["box"] = [170, 120, 20, 25]
    report = preflight(spec, layout_scene(spec, fonts), tmp_path)
    assert any(
        f["code"] == "unintended_overlap" and "vertical" in f["elements"]
        for f in report["findings"]
    )
    page["elements"][-1]["box"] = [95, 140, 20, 20]
    assert preflight(spec, layout_scene(spec, fonts), tmp_path)["status"] == "pass"


def test_native_diagram_styles_survive_pptx_and_svg(scene, fonts, tmp_path):
    spec = diagram(scene)
    layouts = layout_scene(spec, fonts)
    out = tmp_path / "diagram.pptx"
    write_pptx(spec, layouts, tmp_path, out)
    shapes = {s.name: s for s in Presentation(out).slides[0].shapes}
    assert shapes["vertical"].rotation == 90
    assert shapes["vertical"].has_text_frame and shapes["vertical"].text == "Down Sampling"
    assert shapes["arrow"]._element.xpath(".//a:custGeom/a:pathLst/a:path")
    assert shapes["dashed"]._element.xpath(".//a:custDash/a:ds")[0].get("d") == "600000"
    svg = tmp_path / "diagram.svg"
    write_svg(spec["slides"][0], layouts, tmp_path, svg)
    ns = {"s": "http://www.w3.org/2000/svg"}
    xml = ET.parse(svg)
    assert xml.find(".//s:g[@id='vertical']", ns).get("transform") == "rotate(90 180.0 155.0)"
    assert xml.find(".//s:g[@id='dashed']/s:line", ns).get("stroke-dasharray") == "12 8"


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_actual_diagram_render_preserves_orientation_arrow_size_and_dashes(scene, fonts, tmp_path):
    spec = diagram(scene)
    page = spec["slides"][0]
    base = copy.deepcopy(page["elements"][0])
    page["elements"].extend(
        [
            {**base, "id": "negative", "box": [0, 135, 180, 40], "rotation": -90},
            {**base, "id": "upside_down", "box": [280, 290, 180, 40], "rotation": 180},
        ]
    )
    layouts = layout_scene(spec, fonts)
    out = tmp_path / "diagram.pptx"
    write_pptx(spec, layouts, tmp_path, out)
    pdf = make_renderer().render(out, tmp_path / "render")
    report = verify_rendered_text(pdf, spec, layouts)
    assert report["status"] == "pass", report
    png = rasterize_pdf(pdf, tmp_path / "render", width=640)[0]
    im = Image.open(png).convert("RGB")
    red = [
        (x, y)
        for y in range(75, 105)
        for x in range(448, 486)
        if (p := im.getpixel((x, y)))[0] > 170 and p[1] < 120 and p[2] < 120
    ]
    assert 18 <= max(y for x, y in red) - min(y for x, y in red) + 1 <= 22
    row = [
        im.getpixel((x, 230))[2] > 150 and im.getpixel((x, 230))[0] < 120 for x in range(281, 479)
    ]
    transitions = sum(a != b for a, b in zip(row, row[1:], strict=False))
    assert 17 <= transitions <= 21


def test_rotated_container_bounds_and_arrow_empty_space(scene, fonts, tmp_path):
    spec = diagram(scene)
    page = spec["slides"][0]
    text = page["elements"][0]
    text["container"] = "container"
    page["elements"].append(
        {
            "id": "container",
            "kind": "shape",
            "shape": "rect",
            "z": 0,
            "box": [157, 60, 46, 190],
            "fill": "#EEEEEE",
        }
    )
    assert preflight(spec, layout_scene(spec, fonts), tmp_path)["status"] == "pass"
    page["elements"][-1]["box"] = [170, 60, 5, 190]
    report = preflight(spec, layout_scene(spec, fonts), tmp_path)
    assert any(f["code"] == "outside_container" for f in report["findings"])
    del text["container"]
    page["elements"].pop()
    page["elements"].append(
        {
            "id": "clear",
            "kind": "shape",
            "shape": "rect",
            "z": 4,
            "box": [400, 82, 5, 3],
            "fill": "#000000",
        }
    )
    assert preflight(spec, layout_scene(spec, fonts), tmp_path)["status"] == "pass"
    page["elements"][1]["points"] = [[280, 5], [480, 5]]
    report = preflight(spec, layout_scene(spec, fonts), tmp_path)
    assert any(
        f["code"] == "out_of_bounds" and f["elements"] == ["arrow"] for f in report["findings"]
    )


def test_extreme_dash_ratio_is_rejected_before_export(scene):
    spec = diagram(scene)
    spec["slides"][0]["elements"][2]["stroke_width"] = 1e-310
    with pytest.raises(InputError):
        validate_scene(spec)


def test_matching_aspect_stretch_does_not_claim_distortion(scene, fonts, tmp_path):
    spec = diagram(scene)
    Image.new("RGB", (100, 50), "green").save(tmp_path / "image.png")
    spec["slides"][0]["elements"] = [
        {
            "id": "image",
            "kind": "image",
            "z": 0,
            "box": [20, 20, 200, 100],
            "path": "image.png",
            "image_fit": "stretch",
            "contains_text": False,
            "provenance": "Isolated source artwork with matching aspect ratio",
        }
    ]
    assert preflight(spec, layout_scene(spec, fonts), tmp_path)["status"] == "pass"
    spec["slides"][0]["elements"][0]["box"][2] = 210
    report = preflight(spec, layout_scene(spec, fonts), tmp_path)
    assert any(f["code"] == "image_stretched" for f in report["findings"])


@pytest.mark.parametrize("fit", ["contain", "cover", "stretch"])
@pytest.mark.parametrize("rotation", [0, 90])
def test_transparent_artwork_corners_do_not_hide_or_invent_text_collisions(
    scene, fonts, tmp_path, fit, rotation
):
    spec = diagram(scene)
    artwork = Image.new("RGBA", (100, 100))
    ImageDraw.Draw(artwork).rectangle((60, 0, 99, 99), fill="red")
    artwork.save(tmp_path / "artwork.png")
    img = {
        "id": "art",
        "kind": "image",
        "z": 0,
        "box": [0, 0, 120, 80],
        "path": "artwork.png",
        "contains_text": False,
        "image_fit": fit,
        "provenance": "Independent artwork with transparent left side",
    }
    label = {
        "id": "label",
        "kind": "text",
        "z": 1,
        "box": [10, 20, 50, 30],
        "font_family": "Arial",
        "font_size": 20,
        "text": "Hi",
        "rotation": rotation,
    }
    spec["slides"][0]["elements"] = [img, label]
    result = preflight(spec, layout_scene(spec, fonts), tmp_path)
    assert not any(f["severity"] == "error" for f in result["findings"]), result
    label["box"][0] = 78
    result = preflight(spec, layout_scene(spec, fonts), tmp_path)
    assert any(f["code"] == "unintended_overlap" for f in result["findings"])
    label["box"][0] = 10
    img["contains_text"] = True
    result = preflight(spec, layout_scene(spec, fonts), tmp_path)
    assert any(f["code"] == "baked_text_overlap" for f in result["findings"])
