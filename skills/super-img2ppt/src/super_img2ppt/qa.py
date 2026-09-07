"""Geometry and artifact checks, with explicit evidence for each finding."""

from __future__ import annotations

import hashlib
import math
from collections import Counter
from functools import lru_cache
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageStat

from .geometry import box_points, convex_hull, line_parts, polygon_points, rotate_points
from .layout import TextLayout, text_ink_boxes, text_ink_mask
from .scene import InputError, safe_asset, text_content


def polygon(element: dict) -> list[tuple[float, float]]:
    if element["kind"] == "line":
        return convex_hull([p for part in line_parts(element) for p in part])
    x, y, w, h = element["box"]
    shape = element.get("shape")
    if shape == "polygon":
        return polygon_points(element)
    if shape == "ellipse":
        return [
            (
                x + w / 2 + w / 2 * math.cos(i * math.pi / 24),
                y + h / 2 + h / 2 * math.sin(i * math.pi / 24),
            )
            for i in range(48)
        ]
    if shape == "triangle":
        return [(x + w / 2, y), (x + w, y + h), (x, y + h)]
    if shape == "diamond":
        return [(x + w / 2, y), (x + w, y + h / 2), (x + w / 2, y + h), (x, y + h / 2)]
    return rotate_points(box_points(element["box"]), element)


def area(poly: list[tuple[float, float]]) -> float:
    return (
        abs(
            sum(
                poly[i][0] * poly[(i + 1) % len(poly)][1]
                - poly[(i + 1) % len(poly)][0] * poly[i][1]
                for i in range(len(poly))
            )
        )
        / 2
        if poly
        else 0
    )


def intersection(subject: list, clip: list) -> list:
    result = subject
    for index, b in enumerate(clip):
        a = clip[index - 1]
        output = []
        if not result:
            return []

        def side(p, a=a, b=b):
            return (b[0] - a[0]) * (p[1] - a[1]) - (b[1] - a[1]) * (p[0] - a[0])

        prev = result[-1]
        prev_side = side(prev)
        for current in result:
            current_side = side(current)
            if (current_side >= -1e-8) != (prev_side >= -1e-8):
                t = prev_side / (prev_side - current_side)
                output.append(
                    (prev[0] + t * (current[0] - prev[0]), prev[1] + t * (current[1] - prev[1]))
                )
            if current_side >= -1e-8:
                output.append(current)
            prev, prev_side = current, current_side
        result = output
    return result


def is_ancestor(parent: dict, child: dict, elements: dict) -> bool:
    while "container" in child:
        child = elements[child["container"]]
        if child["id"] == parent["id"]:
            return True
    return False


def _alpha_in_text_frame(alpha, image, text, size, left, top):
    x, y, w, h = image["box"]
    iw, ih = alpha.size
    fx, fy = w / iw, h / ih
    fit = image.get("image_fit", "contain")
    if fit != "stretch":
        fx = fy = min(fx, fy) if fit == "contain" else max(fx, fy)
        x += (w - iw * fx) / 2
        y += (h - ih * fy) / 2
    origin, along_x, along_y = rotate_points(
        [(left, top), (left + 0.25, top), (left, top + 0.25)], text
    )
    affine = (
        (along_x[0] - origin[0]) / fx,
        (along_y[0] - origin[0]) / fx,
        (origin[0] - x) / fx,
        (along_x[1] - origin[1]) / fy,
        (along_y[1] - origin[1]) / fy,
        (origin[1] - y) / fy,
    )
    return alpha.transform(
        size, Image.Transform.AFFINE, affine, resample=Image.Resampling.BILINEAR, fillcolor=0
    )


def _text_mask_in_text_frame(other_mask, other, text, size, left, top):
    mask, ox, oy = other_mask
    points = rotate_points([(left, top), (left + 0.25, top), (left, top + 0.25)], text)
    origin, dx, dy = rotate_points(points, other, inverse=True)
    affine = (
        (dx[0] - origin[0]) * 4,
        (dy[0] - origin[0]) * 4,
        (origin[0] - ox) * 4,
        (dx[1] - origin[1]) * 4,
        (dy[1] - origin[1]) * 4,
        (origin[1] - oy) * 4,
    )
    return mask.transform(
        size, Image.Transform.AFFINE, affine, resample=Image.Resampling.BILINEAR, fillcolor=0
    )


def preflight(scene: dict, layouts: dict[tuple[str, str], TextLayout], root: Path) -> dict:
    findings = []
    pages = []

    def issue(slide, code, elements, message, severity="error", **evidence):
        findings.append(
            {
                "slide": slide["id"],
                "code": code,
                "severity": severity,
                "elements": elements,
                "message": message,
                **evidence,
            }
        )

    for slide in scene["slides"]:
        elements = {e["id"]: e for e in slide["elements"]}
        if not slide.get("reviewed", False):
            issue(
                slide,
                "unreviewed_scene",
                [],
                "Confirm OCR and geometry against the source, then set reviewed=true",
                "warning",
            )
        if not elements:
            issue(slide, "empty_slide", [], "Scene has no reconstructed elements")
        source = None
        source_hash = None
        if "source" in slide:
            source = safe_asset(root, slide["source"])
            source_hash = hashlib.sha256(source.read_bytes()).digest()
            with Image.open(source) as im:
                if im.size != (slide["width"], slide["height"]):
                    raise InputError(f"Source dimensions disagree with {slide['id']}")
        polys = {eid: polygon(e) for eid, e in elements.items()}
        footprints = {
            eid: line_parts(e) if e["kind"] == "line" else [polys[eid]]
            for eid, e in elements.items()
        }
        text_footprints = {}
        for eid, element in elements.items():
            if element["kind"] == "text" and layouts[slide["id"], eid].fits:
                ink = text_ink_boxes(element, layouts[slide["id"], eid])
                if ink:
                    text_footprints[eid] = [
                        rotate_points(box_points([x0, y0, x1 - x0, y1 - y0]), element)
                        for x0, y0, x1, y1 in ink
                    ]

        @lru_cache(maxsize=8)
        def ink_mask(eid, elements=elements, slide_id=slide["id"]):
            return text_ink_mask(elements[eid], layouts[slide_id, eid])

        @lru_cache(maxsize=2)
        def artwork_alpha(eid, elements=elements):
            element = elements[eid]
            if element["kind"] != "image" or element["contains_text"]:
                return None
            with Image.open(safe_asset(root, element["path"])) as im:
                if im.width * im.height > 8_000_000:
                    return None
                if "A" not in im.getbands() and "transparency" not in im.info:
                    return None
                return im.convert("RGBA").getchannel("A")

        def ink_area(eid, other, *, outside=False, other_id=None, elements=elements):
            measured = ink_mask(eid)
            if measured is None:
                return None
            mask, left, top = measured
            clip = Image.new("L", mask.size)
            for part in other:
                local = rotate_points(part, elements[eid], inverse=True)
                ImageDraw.Draw(clip).polygon(
                    [((x - left) * 4, (y - top) * 4) for x, y in local], fill=255
                )
            if outside:
                clip = ImageChops.invert(clip)
            if other_id is not None:
                if elements[other_id]["kind"] == "text":
                    other_mask = ink_mask(other_id)
                    if other_mask is None:
                        return None
                    visible = _text_mask_in_text_frame(
                        other_mask, elements[other_id], elements[eid], mask.size, left, top
                    )
                    clip = ImageChops.multiply(clip, visible)
                else:
                    alpha = artwork_alpha(other_id)
                    if alpha is not None:
                        visible = _alpha_in_text_frame(
                            alpha, elements[other_id], elements[eid], mask.size, left, top
                        )
                        clip = ImageChops.multiply(clip, visible)
            coverage = ImageChops.multiply(mask, clip)
            return ImageStat.Stat(coverage).sum[0] / (255 * 16)

        for eid, e in elements.items():
            p = polys[eid]
            # A line's polygon already includes its stroke; shapes have centerline bounds.
            stroke = e.get("stroke_width", 1) / 2 if e.get("stroke") and e["kind"] != "line" else 0
            if (
                min(x for x, _ in p) - stroke < -0.5
                or min(y for _, y in p) - stroke < -0.5
                or max(x for x, _ in p) + stroke > slide["width"] + 0.5
                or max(y for _, y in p) + stroke > slide["height"] + 0.5
            ):
                issue(slide, "out_of_bounds", [eid], "Object extends outside the source page")
            if e.get("confidence", 1) < 0.85:
                issue(
                    slide,
                    "low_confidence",
                    [eid],
                    "Verify this recognition against the source",
                    "warning",
                    confidence=e["confidence"],
                )
            if "container" in e:
                container = polys[e["container"]]
                content = text_footprints.get(eid, [p])
                outside = any(
                    area(intersection(part, container)) < area(part) - 0.5 for part in content
                )
                if outside and eid in text_footprints:
                    refined = ink_area(eid, [container], outside=True)
                    if refined is not None:
                        outside = refined > 0.5
                if outside:
                    issue(
                        slide,
                        "outside_container",
                        [eid, e["container"]],
                        "Visible content is not fully contained by its declared shape",
                    )
                elif area(intersection(p, container)) < area(p) - 0.5:
                    issue(
                        slide,
                        "text_frame_outside_container_only",
                        [eid, e["container"]],
                        "Transparent frame corners extend beyond the container; measured visible ink remains inside",
                        "info",
                    )
            if e["kind"] == "text":
                layout = layouts[slide["id"], eid]
                if not layout.fits:
                    issue(
                        slide,
                        "text_overflow",
                        [eid],
                        "Measured text plus a 0.5 px rendering reserve does not fit; use required_width_px/required_height_px to correct the box or explicitly revise the font",
                        layout=layout.report(),
                    )
                elif layout.scale < 0.9999:
                    issue(
                        slide,
                        "text_shrunk",
                        [eid],
                        "Text was reduced within its configured limit",
                        "info",
                        layout=layout.report(),
                    )
            if e["kind"] == "image":
                path = safe_asset(root, e["path"])
                with Image.open(path) as im:
                    if im.width * im.height > 40_000_000:
                        raise InputError("Asset exceeds 40 million pixels")
                    if im.format not in {"PNG", "JPEG", "WEBP"}:
                        raise InputError("Image assets must be PNG, JPEG or WebP")
                    image_ratio = im.width / im.height
                if source and (
                    path == source or hashlib.sha256(path.read_bytes()).digest() == source_hash
                ):
                    issue(
                        slide,
                        "source_image_as_asset",
                        [eid],
                        "The full source image is a reference, not a reconstructed background",
                    )
                if e["contains_text"]:
                    issue(
                        slide,
                        "raster_text",
                        [eid],
                        "Text inside this independent image remains noneditable",
                        "warning",
                    )
                if e.get("image_fit", "contain") == "stretch" and not math.isclose(
                    e["box"][2] / e["box"][3], image_ratio, rel_tol=0.001
                ):
                    issue(
                        slide,
                        "image_stretched",
                        [eid],
                        "Stretching changes the asset aspect ratio",
                        "warning",
                    )
        ordered = sorted(elements.values(), key=lambda e: e["z"])
        bounds = {
            eid: (
                min(x for x, _ in p),
                min(y for _, y in p),
                max(x for x, _ in p),
                max(y for _, y in p),
            )
            for eid, p in polys.items()
        }
        overlaps = []
        for i, lower in enumerate(ordered):
            for upper in ordered[i + 1 :]:
                lo, hi = lower["id"], upper["id"]
                a, b = bounds[lo], bounds[hi]
                if min(a[2], b[2]) <= max(a[0], b[0]) or min(a[3], b[3]) <= max(a[1], b[1]):
                    continue
                overlap = sum(
                    area(intersection(a, b)) for a in footprints[lo] for b in footprints[hi]
                )
                if overlap <= 0.5:
                    continue
                # Baked text is a separate problem that overlap exemptions cannot waive.
                if {lower["kind"], upper["kind"]} == {"text", "image"}:
                    img = lower if lower["kind"] == "image" else upper
                    if img["contains_text"]:
                        issue(
                            slide,
                            "baked_text_overlap",
                            [lo, hi],
                            "Editable text overlaps an image declared to contain text",
                        )
                        continue
                if (
                    lower.get("role") == "background"
                    and lower["kind"] in {"shape", "image"}
                    and not lower.get("contains_text", False)
                ):
                    continue
                if is_ancestor(lower, upper, elements):
                    continue
                if lo in text_footprints or hi in text_footprints:
                    ink_overlap = sum(
                        area(intersection(a, b))
                        for a in text_footprints.get(lo, footprints[lo])
                        for b in text_footprints.get(hi, footprints[hi])
                    )
                    if ink_overlap > 0.5:
                        text_id, other_id = (lo, hi) if lo in text_footprints else (hi, lo)
                        refined = ink_area(text_id, footprints[other_id], other_id=other_id)
                        if refined is not None:
                            ink_overlap = refined
                    if ink_overlap <= 0.5:
                        issue(
                            slide,
                            "text_frame_overlap_only",
                            [lo, hi],
                            "Text-frame overlap is confined to blank space; measured visible ink remains separate",
                            "info",
                        )
                        continue
                if hi in lower.get("allow_overlap_with", []) or lo in upper.get(
                    "allow_overlap_with", []
                ):
                    overlaps.append(
                        {
                            "elements": [lo, hi],
                            "reason": lower.get("overlap_reason", upper.get("overlap_reason")),
                            "area_px2": round(overlap, 3),
                        }
                    )
                    continue
                issue(
                    slide,
                    "unintended_overlap",
                    [lo, hi],
                    "Objects overlap without a verified container or named exemption",
                    area_px2=round(overlap, 3),
                )
        ink_mask.cache_clear()
        artwork_alpha.cache_clear()
        pages.append(
            {
                "id": slide["id"],
                "counts": dict(Counter(e["kind"] for e in elements.values())),
                "intentional_overlaps": overlaps,
                "text_layouts": {
                    eid: layouts[slide["id"], eid].report()
                    for eid, e in elements.items()
                    if e["kind"] == "text"
                },
            }
        )
    status = (
        "fail"
        if any(f["severity"] == "error" for f in findings)
        else "review"
        if any(f["severity"] == "warning" for f in findings)
        else "pass"
    )
    return {"status": status, "findings": findings, "slides": pages}


def compact_text(text: str) -> str:
    return "".join(text.split())


def inspect_pptx(path: Path, scene: dict) -> dict:
    from pptx import Presentation

    presentation = Presentation(path)
    problems = []
    if len(presentation.slides) != len(scene["slides"]):
        problems.append({"code": "slide_count_mismatch"})
    for index, (slide, expected) in enumerate(
        zip(presentation.slides, scene["slides"], strict=True)
    ):
        shapes = {shape.name: shape for shape in slide.shapes}
        for element in expected["elements"]:
            shape = shapes.get(element["id"])
            if shape is None:
                problems.append(
                    {"slide": index + 1, "element": element["id"], "code": "missing_native_object"}
                )
            elif element["kind"] == "text" and (
                not shape.has_text_frame
                or compact_text(shape.text) != compact_text(text_content(element))
            ):
                problems.append(
                    {
                        "slide": index + 1,
                        "element": element["id"],
                        "code": "text_roundtrip_mismatch",
                    }
                )
        notes = slide.notes_slide.notes_text_frame.text
        if notes != expected.get("notes", ""):
            problems.append({"slide": index + 1, "code": "notes_roundtrip_mismatch"})
    return {
        "status": "fail" if problems else "pass",
        "findings": problems,
        "slide_count": len(presentation.slides),
    }


def comparison(source: Path, preview: Path, out: Path) -> dict:
    with Image.open(source) as raw, Image.open(preview) as rendered:
        src = raw.convert("RGB")
        dst = rendered.convert("RGB").resize(src.size, Image.Resampling.LANCZOS)
    diff = ImageChops.difference(src, dst)
    mae = sum(ImageStat.Stat(diff).mean) / 3 / 255
    heat = diff.convert("L").point(lambda p: min(255, p * 4))
    red = Image.new("RGB", src.size, "#e63946")
    heatmap = Image.composite(red, Image.new("RGB", src.size, "white"), heat)
    sheet = Image.new("RGB", (src.width * 3, src.height + 38), "white")
    draw = ImageDraw.Draw(sheet)
    for i, (label, im) in enumerate(
        [("Source", src), ("Actual PPTX render", dst), ("Difference (4x)", heatmap)]
    ):
        draw.text((i * src.width + 12, 10), label, fill="black")
        sheet.paste(im, (i * src.width, 38))
    sheet.save(out)
    return {
        "mean_absolute_pixel_error": round(mae, 6),
        "comparison": out.name,
        "interpretation": "Diagnostic only; antialiasing and renderer differences affect this number. It is not an acceptance score.",
    }
