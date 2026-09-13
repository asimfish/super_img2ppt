"""Replace explicitly selected raster shapes with native blocks and Office equations."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from jsonschema import Draft202012Validator
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from .prepare import fresh_directory, inspect_archive, json_write
from .render import make_renderer, rasterize_pdf
from .scene import BOX, COLOR, ID, InputError, _unique_pairs

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"


def node(parent, ns, name, **attrs):
    return etree.SubElement(parent, f"{{{ns}}}{name}", attrs)


def validate_expression(expr, depth=0, budget=None):
    budget = [0] if budget is None else budget
    budget[0] += 1
    if depth > 16 or budget[0] > 512:
        raise InputError("Equation exceeds depth/node limits")
    if isinstance(expr, str):
        if not expr or len(expr) > 256 or any(ord(c) < 32 for c in expr):
            raise InputError("Invalid equation token")
    elif isinstance(expr, list):
        if not 1 <= len(expr) <= 128:
            raise InputError("Invalid equation sequence")
        for child in expr:
            validate_expression(child, depth + 1, budget)
    elif isinstance(expr, dict):
        keys = set(expr)
        if keys == {"text", "style"}:
            if expr["style"] not in ("p", "b", "i", "bi"):
                raise InputError("Unknown math style")
            if not isinstance(expr["text"], str):
                raise InputError("Math text must be a string")
            validate_expression(expr["text"], depth + 1, budget)
        elif keys in (
            {"base", "sub"},
            {"base", "sup"},
            {"base", "sub", "sup"},
            {"hat"},
            {"num", "den"},
            {"sqrt"},
        ):
            for child in expr.values():
                validate_expression(child, depth + 1, budget)
        else:
            raise InputError("Unsupported equation structure")
    else:
        raise InputError("Invalid equation node")


def emit_expression(parent, expr, size, color, family):
    if isinstance(expr, list):
        for child in expr:
            emit_expression(parent, child, size, color, family)
        return
    if isinstance(expr, str) or "text" in expr:
        text = expr if isinstance(expr, str) else expr["text"]
        style = (
            ("i" if len(text) == 1 and text.isalpha() else "p")
            if isinstance(expr, str)
            else expr["style"]
        )
        run = node(parent, M, "r")
        props = node(run, M, "rPr")
        if style == "p":
            node(props, M, "nor")
        node(props, M, "sty").set(f"{{{M}}}val", style)
        props = node(run, A, "rPr", sz=str(round(size * 100)), lang="en-US")
        fill = node(props, A, "solidFill")
        node(fill, A, "srgbClr", val=color.lstrip("#"))
        node(props, A, "latin", typeface=family)
        node(run, M, "t").text = text
        return
    if "base" in expr:
        tag = "sSubSup" if "sub" in expr and "sup" in expr else "sSub" if "sub" in expr else "sSup"
        obj = node(parent, M, tag)
        emit_expression(node(obj, M, "e"), expr["base"], size, color, family)
        for key in ("sub", "sup"):
            if key in expr:
                emit_expression(node(obj, M, key), expr[key], size, color, family)
    elif "hat" in expr:
        obj = node(parent, M, "acc")
        props = node(obj, M, "accPr")
        node(props, M, "chr").set(f"{{{M}}}val", "\u0302")
        emit_expression(node(obj, M, "e"), expr["hat"], size, color, family)
    elif "num" in expr:
        obj = node(parent, M, "f")
        for key in ("num", "den"):
            emit_expression(node(obj, M, key), expr[key], size, color, family)
    else:
        obj = node(parent, M, "rad")
        node(node(obj, M, "radPr"), M, "degHide").set(f"{{{M}}}val", "1")
        node(obj, M, "deg")
        emit_expression(node(obj, M, "e"), expr["sqrt"], size, color, family)


def equation(
    slide, box, expression, size, color="#101522", family="STIX Two Math", name="equation"
):
    validate_expression(expression)
    shape = slide.shapes.add_textbox(*map(round, box))
    shape.name = name
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = False
    p = frame.paragraphs[0]._p
    props = node(p, A, "pPr")
    node(props, A, "defRPr", sz=str(round(size * 100)))
    math = node(p, A14, "m")
    para = node(math, M, "oMathPara")
    node(node(para, M, "oMathParaPr"), M, "jc").set(f"{{{M}}}val", "left")
    emit_expression(node(para, M, "oMath"), expression, size, color, family)
    return shape


ITEM = {
    "type": "object",
    "additionalProperties": False,
    "required": ["id", "kind", "box"],
    "properties": {
        "id": ID,
        "kind": {"enum": ["cuboid", "equation", "text"]},
        "box": BOX,
        "expression": {},
        "text": {"type": "string", "minLength": 1, "maxLength": 256},
        "font_size": {"type": "number", "minimum": 1, "maximum": 200},
        "color": COLOR,
        "front": COLOR,
        "top": COLOR,
        "side": COLOR,
        "depth": {
            "type": "array",
            "items": {"type": "number", "minimum": 0, "maximum": 500},
            "minItems": 2,
            "maxItems": 2,
        },
    },
    "allOf": [
        {
            "if": {"properties": {"kind": {"const": "cuboid"}}},
            "then": {"required": ["depth", "front", "top", "side"]},
        },
        {
            "if": {"properties": {"kind": {"const": "equation"}}},
            "then": {"required": ["expression", "font_size", "color"]},
        },
        {
            "if": {"properties": {"kind": {"const": "text"}}},
            "then": {"required": ["text", "font_size", "color"]},
        },
    ],
}
SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["version", "source_sha256", "width", "height", "replacements"],
    "properties": {
        "version": {"const": 1},
        "source_sha256": {"type": "string", "pattern": "^[0-9a-f]{64}$"},
        "width": {"type": "integer", "minimum": 16, "maximum": 16384},
        "height": {"type": "integer", "minimum": 16, "maximum": 16384},
        "replacements": {
            "type": "array",
            "minItems": 1,
            "maxItems": 200,
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["slide", "image", "items"],
                "properties": {
                    "slide": {"type": "integer", "minimum": 1, "maximum": 200},
                    "image": ID,
                    "items": {"type": "array", "minItems": 1, "maxItems": 256, "items": ITEM},
                },
            },
        },
    },
}


def polygon(slide, pts, color, name):
    builder = slide.shapes.build_freeform(*map(round, pts[0]))
    builder.add_line_segments([tuple(map(round, p)) for p in pts[1:]], close=True)
    shape = builder.convert_to_shape()
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color[1:])
    shape.line.fill.background()
    # Freeforms inherit a theme effect otherwise (visible drop shadows in LibreOffice).
    node(shape._element.spPr, A, "effectLst")
    for effect in shape._element.xpath("./p:style/a:effectRef"):
        effect.set("idx", "0")
    return shape


def walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def replace_rasters(source: Path, plan: dict, out: Path, render=True):
    fresh_directory(out)
    report = {"status": "fail", "visual_review": "required", "replacements": []}
    try:
        if list(Draft202012Validator(SCHEMA).iter_errors(plan)):
            raise InputError("Invalid native replacement plan")
        # Reject NaN even in arbitrary expression input before native parsers run.
        json.dumps(plan, allow_nan=False)
        if source.stat().st_size > 200_000_000:
            raise InputError("PPTX exceeds 200 MB")
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        if digest != plan["source_sha256"]:
            raise InputError("Native replacement source hash mismatch")
        inspect_archive(source)
        deck = Presentation(source)
        scale_x, scale_y = deck.slide_width / plan["width"], deck.slide_height / plan["height"]
        if abs(scale_x / scale_y - 1) > 0.001:
            raise InputError("Replacement canvas must match slide aspect ratio")
        existing_names = {shape.name for sl in deck.slides for shape in walk_shapes(sl.shapes)}
        seen, item_ids, count = set(), set(), 0
        for spec in plan["replacements"]:
            key = spec["slide"], spec["image"]
            if key in seen or spec["slide"] > len(deck.slides):
                raise InputError("Duplicate or invalid replacement target")
            seen.add(key)
            slide = deck.slides[spec["slide"] - 1]
            matches = [s for s in slide.shapes if s.name == spec["image"]]
            if len(matches) != 1 or matches[0].shape_type != MSO_SHAPE_TYPE.PICTURE:
                raise InputError("Replacement requires one top-level picture with the exact name")
            original = matches[0]
            tree = original._element.getparent()
            offset = tree.index(original._element)
            for item in spec["items"]:
                count += 1
                if count > 1500 or item["id"] in item_ids:
                    raise InputError("Too many or duplicate native items")
                if item["id"] in existing_names and item["id"] != original.name:
                    raise InputError("Native item name conflicts with an existing shape")
                item_ids.add(item["id"])
                x, y, w, h = item["box"]
                if x < 0 or y < 0 or x + w > plan["width"] or y + h > plan["height"]:
                    raise InputError("Native item outside canvas")
                box = [x * scale_x, y * scale_y, w * scale_x, h * scale_y]
                size = item.get("font_size", 18) * scale_x / 12700
                if item["kind"] == "equation":
                    shape = equation(
                        slide, box, item["expression"], size, item["color"], name=item["id"]
                    )
                elif item["kind"] == "text":
                    shape = slide.shapes.add_textbox(*map(round, box))
                    shape.name = item["id"]
                    tf = shape.text_frame
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.word_wrap = False
                    p = tf.paragraphs[0]
                    p.text = item["text"]
                    p.font.name = "Times New Roman"
                    p.font.size = Pt(size)
                    p.font.color.rgb = RGBColor.from_string(item["color"][1:])
                else:
                    dx, dy = item["depth"]
                    if dx <= 0 or dy <= 0 or dx >= w or dy >= h:
                        raise InputError("Cuboid depth must fit its bounding box")
                    faces = [
                        (
                            [(x, y + dy), (x + w - dx, y + dy), (x + w - dx, y + h), (x, y + h)],
                            "front",
                        ),
                        ([(x, y + dy), (x + dx, y), (x + w, y), (x + w - dx, y + dy)], "top"),
                        (
                            [
                                (x + w - dx, y + dy),
                                (x + w, y),
                                (x + w, y + h - dy),
                                (x + w - dx, y + h),
                            ],
                            "side",
                        ),
                    ]
                    shapes = [
                        polygon(
                            slide,
                            [(a * scale_x, b * scale_y) for a, b in pts],
                            item[face],
                            item["id"] + "_" + face,
                        )
                        for pts, face in faces
                    ]
                    shape = slide.shapes.add_group_shape(shapes)
                    shape.name = item["id"]
                tree.remove(shape._element)
                tree.insert(offset, shape._element)
                offset += 1
            tree.remove(original._element)
            report["replacements"].append(
                {"slide": spec["slide"], "image": spec["image"], "native_items": len(spec["items"])}
            )
        output = out / "editable.pptx"
        deck.save(output)
        report.update(
            native_equations=sum(
                i["kind"] == "equation" for r in plan["replacements"] for i in r["items"]
            ),
            native_cuboid_groups=sum(
                i["kind"] == "cuboid" for r in plan["replacements"] for i in r["items"]
            ),
            math_font="STIX Two Math (not embedded)",
            source_sha256=digest,
            output_sha256=hashlib.sha256(output.read_bytes()).hexdigest(),
            remaining_images=[
                {"slide": i + 1, "name": s.name}
                for i, sl in enumerate(deck.slides)
                for s in walk_shapes(sl.shapes)
                if s.shape_type == MSO_SHAPE_TYPE.PICTURE
            ],
            limitations=[
                "PPTX only; original scene/SVG are not modified. Replay this plan after rebuilding the source.",
                "Native Office equations; PowerPoint/WPS editing not tested by LibreOffice rendering.",
                "Replaced regions require source-vs-render visual and color review; no full-fidelity pass.",
            ],
        )
        json_write(out / "plan.json", plan)
        if render:
            pdf = make_renderer().render(output, out / "render")
            rasterize_pdf(pdf, out / "render", width=plan["width"])
        report["status"] = "review" if render else "unverified"
        return report
    finally:
        json_write(out / "native.json", report)


def load_plan(path):
    if path.stat().st_size > 2_000_000:
        raise InputError("Native replacement plan exceeds 2 MB")
    return json.loads(path.read_text(), object_pairs_hook=_unique_pairs)
