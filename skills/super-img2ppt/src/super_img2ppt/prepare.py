"""Normalize user-selected inputs and retain source evidence and speaker notes."""

from __future__ import annotations

import hashlib
import json
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET

from PIL import Image, ImageOps
from pptx import Presentation

from .ocr import choose_backend
from .render import make_renderer, rasterize_pdf
from .scene import InputError

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".tiff", ".tif", ".bmp"}


def json_write(path: Path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def fresh_directory(path: Path):
    if path.exists() or path.is_symlink():
        raise InputError(f"Output directory already exists; use a new run directory: {path}")
    path.mkdir(parents=True)


def inspect_archive(path: Path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 20000 or sum(e.file_size for e in entries) > 800_000_000:
            raise InputError("PPTX archive exceeds resource limits")
        names = [e.filename for e in entries]
        if len(names) != len(set(names)):
            raise InputError("PPTX contains duplicate archive members")
        for entry in entries:
            member = PurePosixPath(entry.filename)
            if member.is_absolute() or ".." in member.parts or "\\" in entry.filename:
                raise InputError("Unsafe archive member path")
            if (
                entry.file_size > 200_000_000
                or entry.file_size / max(1, entry.compress_size) > 1000
            ):
                raise InputError("PPTX archive member exceeds resource limits")
            if "vbaproject" in entry.filename.lower() or "/embeddings/" in entry.filename.lower():
                raise InputError("Macro and embedded executable/object parts are not supported")
            if entry.filename.endswith((".rels", ".xml")):
                if entry.file_size > 16_000_000:
                    raise InputError("PPTX XML part exceeds 16 MB")
                data = archive.read(entry)
                if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
                    raise InputError("XML declarations/entities are not permitted")
                if entry.filename.endswith(".rels"):
                    xml = ET.fromstring(data)
                    for node in xml:
                        if node.attrib.get("TargetMode", "").lower() == "external":
                            raise InputError(
                                "PPTX external relationships require a reviewed local copy with external links removed"
                            )


def _normalize_image(source: Path, dest: Path):
    with Image.open(source) as raw:
        if raw.width * raw.height > 40_000_000:
            raise InputError("Source image exceeds 40 million pixels")
        if getattr(raw, "n_frames", 1) > 1:
            raise InputError("Multi-frame images must be split into explicit pages first")
        im = ImageOps.exif_transpose(raw)
        if im.width < 16 or im.height < 16 or max(im.size) > 16384:
            raise InputError("Source dimensions must be between 16 and 16384 pixels")
        rgba = im.convert("RGBA")
        background = Image.new("RGBA", im.size, "white")
        background.alpha_composite(rgba)
        background.convert("RGB").save(dest)


def _inputs(paths: list[Path]) -> list[Path]:
    result = []
    for path in paths:
        if path.is_dir():

            def natural_key(p):
                return [
                    (0, int(s)) if s.isdigit() else (1, s.casefold())
                    for s in re.split(r"(\d+)", p.name)
                ]

            result.extend(
                sorted(
                    (
                        p
                        for p in path.iterdir()
                        if p.is_file() and p.suffix.lower() in IMAGE_SUFFIXES
                    ),
                    key=natural_key,
                )
            )
        else:
            result.append(path)
    if not result or len(result) > 200:
        raise InputError("Supply between 1 and 200 image/PDF/PPTX files")
    for path in result:
        if not path.is_file() or path.stat().st_size > 200_000_000:
            raise InputError(f"Input is not a regular file under 200 MB: {path}")
        if path.suffix.lower() not in IMAGE_SUFFIXES | {".pdf", ".pptx"}:
            raise InputError(f"Unsupported input type: {path.suffix}")
    return result


def prepare(paths: list[Path], out: Path, ocr: str = "auto", languages: str | None = None) -> dict:
    inputs = _inputs(paths)
    fresh_directory(out)
    backend_name, backend = choose_backend(ocr, languages)
    slides = []
    sources = []
    for source in inputs:
        sources.append(
            {"name": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest()}
        )
        with tempfile.TemporaryDirectory(prefix="super-img2ppt-prepare-") as tmp:
            temp = Path(tmp)
            notes = []
            if source.suffix.lower() == ".pptx":
                inspect_archive(source)
                original = Presentation(source)
                notes = [
                    s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
                    for s in original.slides
                ]
                pdf = make_renderer().render(source, temp)
                pages = rasterize_pdf(pdf, temp)
                if len(notes) != len(pages):
                    raise InputError(
                        "PPTX renderer changed page count; notes mapping would be ambiguous"
                    )
            elif source.suffix.lower() == ".pdf":
                pages = rasterize_pdf(source, temp)
            else:
                pages = [source]
            if len(slides) + len(pages) > 200:
                raise InputError("Combined input exceeds 200 pages")
            for index, page in enumerate(pages):
                page_id = f"page_{len(slides) + 1:03d}"
                page_dir = out / "pages" / page_id
                page_dir.mkdir(parents=True)
                normalized = page_dir / "source.png"
                _normalize_image(page, normalized)
                with Image.open(normalized) as im:
                    width, height = im.size
                hints = {
                    "backend": backend_name,
                    "status": "unavailable" if backend is None else "ok",
                    "coordinate_system": "source pixels; top-left origin; boxes bound ink, not font em or paragraph height",
                    "lines": [],
                }
                if backend is not None:
                    try:
                        hints["lines"] = backend.recognize(normalized)
                    except (InputError, OSError, subprocess.TimeoutExpired) as exc:
                        hints.update(status="failed", error=str(exc))
                json_write(page_dir / "ocr.json", hints)
                slides.append(
                    {
                        "id": page_id,
                        "width": width,
                        "height": height,
                        "background": "#FFFFFF",
                        "source": f"pages/{page_id}/source.png",
                        "notes": notes[index] if notes else "",
                        "reviewed": False,
                        "elements": [],
                    }
                )
    scene = {"version": 1, "title": inputs[0].stem, "slides": slides}
    json_write(out / "scene.json", scene)
    manifest = {
        "sources": sources,
        "page_count": len(slides),
        "ocr": backend_name,
        "status": "needs_reconstruction",
        "scene": "scene.json",
        "next": "Inspect each source image and OCR hints, reconstruct native objects in scene.json, then build. Empty drafts cannot pass QA.",
    }
    json_write(out / "prepare.json", manifest)
    return manifest
