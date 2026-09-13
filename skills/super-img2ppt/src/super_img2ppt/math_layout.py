"""Calibrate native equation frames against LibreOffice's actual math import.

LibreOffice stretches imported OMML to the text shape's box. The ODP object's
intrinsic view-area ratio and the PDF's measured font size expose that stretch.
This does not change equation tokens or certify other Office viewers.
"""

import math
import re
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile

import pypdfium2 as pdfium
from lxml import etree
from pptx import Presentation

from .native import M
from .prepare import inspect_archive
from .scene import InputError

D = "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
C = "urn:oasis:names:tc:opendocument:xmlns:config:1.0"
X = "http://www.w3.org/1999/xlink"


def intrinsic_ratios(odp, targets):
    inspect_archive(odp)
    result = {}
    with ZipFile(odp) as archive:
        root = etree.fromstring(archive.read("content.xml"))
        pages = root.findall(f".//{{{D}}}page")
        for index, page in enumerate(pages):
            for frame in page.findall(f".//{{{D}}}frame"):
                key = index + 1, frame.get(f"{{{D}}}name")
                if key not in targets:
                    continue
                obj = frame.find(f"{{{D}}}object")
                if obj is None:
                    continue
                href = obj.get(f"{{{X}}}href", "").removeprefix("./")
                if not re.fullmatch(r"Object \d+", href):
                    raise InputError("Unexpected native equation object reference")
                try:
                    settings = etree.fromstring(archive.read(href + "/settings.xml"))
                    vals = {e.get(f"{{{C}}}name"): e.text for e in settings.iter()}
                    width, height = float(vals["ViewAreaWidth"]), float(vals["ViewAreaHeight"])
                except (KeyError, TypeError, ValueError) as exc:
                    raise InputError("Missing or invalid native equation metrics") from exc
                if not all(math.isfinite(v) and v > 0 for v in (width, height)):
                    raise InputError("Invalid intrinsic equation size")
                key = index + 1, frame.get(f"{{{D}}}name")
                if key in result:
                    raise InputError("Duplicate equation names in rendered document")
                result[key] = width / height
    return result


def fit_box(box, ratio, measured_pt, target_pt):
    if not all(math.isfinite(v) and v > 0 for v in (ratio, measured_pt, target_pt)):
        raise InputError("Invalid equation size measurement")
    x, y, width, height = box
    new_height = height * target_pt / measured_pt
    new_width = new_height * ratio
    return [x + (width - new_width) / 2, y + (height - new_height) / 2, new_width, new_height]


def calibrate_equations(pptx, pdf, targets, out, renderer):
    """Targets map (1-based slide, top-level equation name) to desired point sizes."""
    out.mkdir()
    with tempfile.TemporaryDirectory(prefix="super-img2ppt-math-") as tmp:
        args = [
            renderer.executable,
            f"-env:UserInstallation={Path(tmp).as_uri()}",
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--norestore",
            "--convert-to",
            "odp",
            "--outdir",
            str(out.resolve()),
            str(pptx.resolve()),
        ]
        try:
            result = subprocess.run(
                args, capture_output=True, text=True, timeout=renderer.timeout, check=False
            )
        except subprocess.TimeoutExpired as exc:
            raise InputError("Native math measurement timed out") from exc
    odp = out / (pptx.stem + ".odp")
    if result.returncode != 0 or not odp.exists():
        raise InputError(
            "Native math measurement failed: " + (result.stderr + result.stdout)[-1000:]
        )
    ratios = intrinsic_ratios(odp, targets)
    deck = Presentation(pptx)
    rows = []
    with pdfium.PdfDocument(pdf) as document:
        if len(document) != len(deck.slides):
            raise InputError("Equation PDF slide count differs")
        for index, slide in enumerate(deck.slides):
            page = document[index]
            text = page.get_textpage()
            try:
                glyphs = []
                for i in range(text.count_chars()):
                    if not text.get_text_range(i, 1).strip():
                        continue
                    a, b, c, d = text.get_charbox(i)
                    size = pdfium.raw.FPDFText_GetFontSize(text, i)
                    glyphs.append(((a + c) / 2, page.get_height() - (b + d) / 2, size))
                for shape in slide.shapes:
                    key = index + 1, shape.name
                    if key not in targets:
                        continue
                    if not shape._element.findall(f".//{{{M}}}oMath") or key not in ratios:
                        raise InputError("Native equation measurement missing")
                    old = [v / 12700 for v in (shape.left, shape.top, shape.width, shape.height)]
                    x, y, w, h = old
                    sizes = [
                        s for gx, gy, s in glyphs if x <= gx <= x + w and y <= gy <= y + h and s > 0
                    ]
                    if not sizes:
                        raise InputError("Equation has no measurable glyphs")
                    measured = max(sizes)
                    new = fit_box(old, ratios[key], measured, targets[key])
                    nx, ny, nw, nh = new
                    if (
                        nx < 0
                        or ny < 0
                        or nx + nw > deck.slide_width / 12700
                        or ny + nh > deck.slide_height / 12700
                    ):
                        raise InputError("Calibrated equation exceeds canvas; revise its placement")
                    shape.left, shape.top, shape.width, shape.height = [
                        round(v * 12700) for v in new
                    ]
                    rows.append(
                        dict(
                            slide=index + 1,
                            name=shape.name,
                            before_box_pt=old,
                            after_box_pt=new,
                            measured_font_pt=measured,
                            target_font_pt=targets[key],
                            intrinsic_aspect=ratios[key],
                        )
                    )
            finally:
                text.close()
                page.close()
    if len(rows) != len(targets):
        raise InputError("Not every target equation was calibrated")
    deck.save(pptx)
    return {
        "status": "review",
        "equations": rows,
        "scope": "One-pass LibreOffice frame calibration; source alignment and other viewers require review",
        "measurement_limit": "Maximum glyph size in the equation frame can include overlapping text or enlarged symbols; inspect final render and remeasure isolated regions.",
    }
