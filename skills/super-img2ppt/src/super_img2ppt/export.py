"""PPTX and SVG exporters sharing one measured scene."""

from __future__ import annotations

import base64
import io
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from .fonts import east_asian_name
from .geometry import arrow_outline, polygon_points
from .layout import TextLayout, positioned_lines
from .scene import safe_asset, slide_transform

SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "round_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "chevron": MSO_SHAPE.CHEVRON,
}


def _color(color: str) -> RGBColor:
    return RGBColor.from_string(color.removeprefix("#"))


def _set_font(font, fragment, point_scale):
    font.name = east_asian_name(fragment.face)
    font.size = Pt(fragment.size * point_scale)
    font.bold = fragment.bold
    font.italic = fragment.italic
    font.color.rgb = _color(fragment.color)
    properties = font._rPr
    properties.set("lang", "zh-CN")
    properties.set("dirty", "0")
    # Avoid Office theme substitution for East Asian and complex-script glyphs.
    for tag in ["a:ea", "a:cs"]:
        child = properties.find(
            f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag[2:]}"
        )
        if child is None:
            child = OxmlElement(tag)
            properties.append(child)
        child.set("typeface", east_asian_name(fragment.face))


def _fill_and_line(shape, element: dict, point_scale: float):
    if "gradient" in element:
        shape.fill.gradient()
        grad = shape._element.spPr.gradFill
        for child in list(grad):
            grad.remove(child)
        stops = OxmlElement("a:gsLst")
        for stop in element["gradient"]["stops"]:
            gs = OxmlElement("a:gs")
            gs.set("pos", str(round(stop["offset"] * 100000)))
            color = OxmlElement("a:srgbClr")
            color.set("val", stop["color"].removeprefix("#"))
            gs.append(color)
            stops.append(gs)
        grad.append(stops)
        linear = OxmlElement("a:lin")
        linear.set("ang", "5400000" if element["gradient"]["direction"] == "vertical" else "0")
        linear.set("scaled", "1")
        grad.append(linear)
    elif element.get("fill") is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _color(element["fill"])
    else:
        shape.fill.background()
    if element.get("stroke") is not None:
        shape.line.color.rgb = _color(element["stroke"])
        shape.line.width = Pt(element.get("stroke_width", 1) * point_scale)
        _dash(shape, element)
    else:
        shape.line.fill.background()


def _dash(shape, element):
    if "dash" not in element:
        return
    dash = OxmlElement("a:custDash")
    pair = OxmlElement("a:ds")
    on, off = element["dash"]
    width = element.get("stroke_width", 1)
    pair.set("d", str(round(on / width * 100000)))
    pair.set("sp", str(round(off / width * 100000)))
    dash.append(pair)
    shape.line._get_or_add_ln().append(dash)


def _image_data(path: Path) -> tuple[io.BytesIO, tuple[int, int]]:
    with Image.open(path) as im:
        data = io.BytesIO()
        im.save(data, format="PNG")
        data.seek(0)
        return data, im.size


def write_pptx(scene: dict, layouts: dict[tuple[str, str], TextLayout], root: Path, out: Path):
    deck = Presentation()
    transform = slide_transform(scene, scene["slides"][0])
    deck.slide_width = Inches(transform.width_inches)
    deck.slide_height = Inches(transform.height_inches)
    deck.core_properties.title = scene.get("title", "Editable reconstruction")
    deck.core_properties.subject = (
        "Reconstructed from source images; see validation and font manifest"
    )
    for spec in scene["slides"]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _color(spec.get("background", "#FFFFFF"))
        tx = slide_transform(scene, spec)
        point_scale = tx.scale * 72
        for element in sorted(spec["elements"], key=lambda e: e["z"]):
            kind = element["kind"]
            if kind == "line":
                if "arrow_head" in element:
                    points = [(x * 1000, y * 1000) for x, y in arrow_outline(element)]
                    builder = slide.shapes.build_freeform(*points[0], scale=tx.scale * 914.4)
                    builder.add_line_segments(points[1:], close=True)
                    shape = builder.convert_to_shape(Inches(tx.x), Inches(tx.y))
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = _color(element.get("stroke", "#111827"))
                    shape.line.fill.background()
                else:
                    p1, p2 = [tx.point(*point) for point in element["points"]]
                    shape = slide.shapes.add_connector(
                        MSO_CONNECTOR.STRAIGHT,
                        Inches(p1[0]),
                        Inches(p1[1]),
                        Inches(p2[0]),
                        Inches(p2[1]),
                    )
                    shape.line.color.rgb = _color(element.get("stroke", "#111827"))
                    shape.line.width = Pt(element.get("stroke_width", 1) * point_scale)
                    _dash(shape, element)
                    if element.get("arrow", False):
                        arrow = OxmlElement("a:tailEnd")
                        arrow.set("type", "triangle")
                        shape.line._get_or_add_ln().append(arrow)
            else:
                x, y, w, h = tx.box(element["box"])
                if kind == "text":
                    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
                    shape.rotation = element.get("rotation", 0)
                    frame = shape.text_frame
                    frame.clear()
                    frame.word_wrap = False
                    frame.auto_size = MSO_AUTO_SIZE.NONE
                    frame.vertical_anchor = {
                        "top": MSO_ANCHOR.TOP,
                        "middle": MSO_ANCHOR.MIDDLE,
                        "bottom": MSO_ANCHOR.BOTTOM,
                    }[element.get("valign", "top")]
                    top, right, bottom, left = element.get("padding", [0, 0, 0, 0])
                    frame.margin_top, frame.margin_right, frame.margin_bottom, frame.margin_left = [
                        Inches(v * tx.scale) for v in [top, right, bottom, left]
                    ]
                    frame._txBody.bodyPr.set("anchorCtr", "0")
                    frame._txBody.bodyPr.set("vertOverflow", "overflow")
                    layout = layouts[spec["id"], element["id"]]
                    for index, line in enumerate(layout.lines):
                        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                        p.alignment = {
                            "left": PP_ALIGN.LEFT,
                            "center": PP_ALIGN.CENTER,
                            "right": PP_ALIGN.RIGHT,
                        }[element.get("align", "left")]
                        p.space_before = Pt(0)
                        p.space_after = Pt(0)
                        p.line_spacing = Pt(line.height * point_scale)
                        p._p.get_or_add_pPr().set("fontAlgn", "base")
                        _set_font(p.font, line.fragments[0], point_scale)
                        for fragment in line.fragments:
                            run = p.add_run()
                            run.text = fragment.text
                            _set_font(run.font, fragment, point_scale)
                elif kind == "shape":
                    if element["shape"] == "polygon":
                        points = [(px * 1000, py * 1000) for px, py in polygon_points(element)]
                        builder = slide.shapes.build_freeform(*points[0], scale=tx.scale * 914.4)
                        builder.add_line_segments(points[1:], close=True)
                        shape = builder.convert_to_shape(Inches(tx.x), Inches(tx.y))
                    else:
                        shape = slide.shapes.add_shape(
                            SHAPES[element["shape"]], Inches(x), Inches(y), Inches(w), Inches(h)
                        )
                    _fill_and_line(shape, element, point_scale)
                    if element["shape"] == "polygon":
                        shape.line._get_or_add_ln().append(OxmlElement("a:round"))
                    if element["shape"] == "round_rect":
                        radius = element.get("radius", min(element["box"][2:]) * 0.12)
                        shape.adjustments[0] = min(0.5, radius / min(element["box"][2:]))
                    elif element["shape"] == "chevron":
                        shape.adjustments[0] = 0.25
                elif kind == "image":
                    data, (iw, ih) = _image_data(safe_asset(root, element["path"]))
                    fit = element.get("image_fit", "contain")
                    if fit == "contain":
                        factor = min(w / iw, h / ih)
                        dw, dh = iw * factor, ih * factor
                        shape = slide.shapes.add_picture(
                            data,
                            Inches(x + (w - dw) / 2),
                            Inches(y + (h - dh) / 2),
                            Inches(dw),
                            Inches(dh),
                        )
                    else:
                        shape = slide.shapes.add_picture(
                            data, Inches(x), Inches(y), Inches(w), Inches(h)
                        )
                        if fit == "cover":
                            scale = max(w / iw, h / ih)
                            cx = (1 - w / (iw * scale)) / 2
                            cy = (1 - h / (ih * scale)) / 2
                            shape.crop_left = shape.crop_right = cx
                            shape.crop_top = shape.crop_bottom = cy
            shape.name = element["id"]
            shape.shadow.inherit = False
            for effect_ref in shape._element.xpath("./p:style/a:effectRef"):
                effect_ref.set("idx", "0")
            cnv = shape._element.xpath(".//p:cNvPr")[0]
            cnv.set(
                "descr", element.get("provenance", f"Editable {kind}; source id {element['id']}")
            )
        slide.notes_slide.notes_text_frame.text = spec.get("notes", "")
    deck.save(out)


def write_svg(slide: dict, layouts: dict[tuple[str, str], TextLayout], root: Path, out: Path):
    svg = ET.Element(
        "svg",
        {
            "xmlns": "http://www.w3.org/2000/svg",
            "viewBox": f"0 0 {slide['width']} {slide['height']}",
            "width": str(slide["width"]),
            "height": str(slide["height"]),
        },
    )
    ET.SubElement(
        svg, "rect", {"width": "100%", "height": "100%", "fill": slide.get("background", "#FFFFFF")}
    )
    for element in sorted(slide["elements"], key=lambda e: e["z"]):
        group = ET.SubElement(svg, "g", {"id": element["id"]})
        kind = element["kind"]
        if kind == "text":
            if element.get("rotation", 0):
                x, y, w, h = element["box"]
                group.set("transform", f"rotate({element['rotation']} {x + w / 2} {y + h / 2})")
            for x, y, line in positioned_lines(element, layouts[slide["id"], element["id"]]):
                text_node = ET.SubElement(
                    group,
                    "text",
                    {
                        "x": str(x),
                        "y": str(y + line.ascent),
                        "{http://www.w3.org/XML/1998/namespace}space": "preserve",
                    },
                )
                for fragment in line.fragments:
                    tspan = ET.SubElement(
                        text_node,
                        "tspan",
                        {
                            "font-family": fragment.face.family,
                            "font-size": str(fragment.size),
                            "font-weight": "bold" if fragment.bold else "normal",
                            "font-style": "italic" if fragment.italic else "normal",
                            "fill": fragment.color,
                        },
                    )
                    tspan.text = fragment.text
        elif kind == "line":
            if "arrow_head" in element:
                ET.SubElement(
                    group,
                    "polygon",
                    {
                        "points": " ".join(f"{x},{y}" for x, y in arrow_outline(element)),
                        "fill": element.get("stroke", "#111827"),
                    },
                )
                continue
            (x1, y1), (x2, y2) = element["points"]
            attrs = {
                "x1": str(x1),
                "y1": str(y1),
                "x2": str(x2),
                "y2": str(y2),
                "stroke": element.get("stroke", "#111827"),
                "stroke-width": str(element.get("stroke_width", 1)),
            }
            if "dash" in element:
                attrs["stroke-dasharray"] = " ".join(map(str, element["dash"]))
            if element.get("arrow", False):
                marker_id = f"marker-{element['id']}"
                marker = ET.SubElement(
                    ET.SubElement(group, "defs"),
                    "marker",
                    {
                        "id": marker_id,
                        "markerWidth": "6",
                        "markerHeight": "6",
                        "refX": "6",
                        "refY": "3",
                        "orient": "auto",
                    },
                )
                ET.SubElement(marker, "path", {"d": "M0,0 L6,3 L0,6 Z", "fill": attrs["stroke"]})
                attrs["marker-end"] = f"url(#{marker_id})"
            ET.SubElement(group, "line", attrs)
        else:
            x, y, w, h = element["box"]
            attrs = {"x": str(x), "y": str(y), "width": str(w), "height": str(h)}
            if kind == "image":
                data, _ = _image_data(safe_asset(root, element["path"]))
                attrs["href"] = "data:image/png;base64," + base64.b64encode(data.getvalue()).decode(
                    "ascii"
                )
                attrs["preserveAspectRatio"] = {
                    "contain": "xMidYMid meet",
                    "cover": "xMidYMid slice",
                    "stretch": "none",
                }[element.get("image_fit", "contain")]
                ET.SubElement(group, "image", attrs)
            else:
                style = {
                    "fill": element.get("fill") or "none",
                    "stroke": element.get("stroke") or "none",
                    "stroke-width": str(element.get("stroke_width", 1)),
                }
                if "dash" in element:
                    style["stroke-dasharray"] = " ".join(map(str, element["dash"]))
                if "gradient" in element:
                    gid = f"gradient-{element['id']}"
                    vertical = element["gradient"]["direction"] == "vertical"
                    gradient = ET.SubElement(
                        ET.SubElement(group, "defs"),
                        "linearGradient",
                        {
                            "id": gid,
                            "x1": "0%",
                            "y1": "0%",
                            "x2": "0%" if vertical else "100%",
                            "y2": "100%" if vertical else "0%",
                        },
                    )
                    for stop in element["gradient"]["stops"]:
                        ET.SubElement(
                            gradient,
                            "stop",
                            {"offset": str(stop["offset"]), "stop-color": stop["color"]},
                        )
                    style["fill"] = f"url(#{gid})"
                name = element["shape"]
                if name in {"rect", "round_rect"}:
                    if name == "round_rect":
                        attrs["rx"] = str(
                            min(min(w, h) / 2, element.get("radius", min(w, h) * 0.12))
                        )
                    ET.SubElement(group, "rect", {**attrs, **style})
                elif name == "ellipse":
                    ET.SubElement(
                        group,
                        "ellipse",
                        {
                            "cx": str(x + w / 2),
                            "cy": str(y + h / 2),
                            "rx": str(w / 2),
                            "ry": str(h / 2),
                            **style,
                        },
                    )
                else:
                    if name == "polygon":
                        points = polygon_points(element)
                        style["stroke-linejoin"] = "round"
                    elif name == "triangle":
                        points = [(x + w / 2, y), (x + w, y + h), (x, y + h)]
                    elif name == "diamond":
                        points = [
                            (x + w / 2, y),
                            (x + w, y + h / 2),
                            (x + w / 2, y + h),
                            (x, y + h / 2),
                        ]
                    else:
                        offset = min(w, h) * 0.25
                        points = [
                            (x, y),
                            (x + w - offset, y),
                            (x + w, y + h / 2),
                            (x + w - offset, y + h),
                            (x, y + h),
                            (x + offset, y + h / 2),
                        ]
                    ET.SubElement(
                        group,
                        "polygon",
                        {"points": " ".join(f"{px},{py}" for px, py in points), **style},
                    )
    ET.indent(svg)
    ET.ElementTree(svg).write(out, encoding="utf-8", xml_declaration=True)
